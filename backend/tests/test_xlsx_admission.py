from __future__ import annotations

import asyncio
import tracemalloc
import zipfile
from io import BytesIO

import pytest
from openpyxl import Workbook
from starlette.datastructures import Headers, UploadFile

from app import ingestion_budget
from app.ingestion_budget import (
    CELL_COST_BYTES,
    DEFERRED_ROW_COST_BYTES,
    ROW_COST_BYTES,
    IngestionBudget,
)
from app.upload_security import receive_and_inspect
from app.xlsx_package import (
    READ_BLOCK,
    ByteLedger,
    PackageViolation,
    ScanningReader,
)

XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
SHEET_NS = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
pytestmark = pytest.mark.security


def _workbook_bytes(
    *, rows: int = 2, columns: int = 2, formulas: int = 0, extra: dict | None = None
) -> bytes:
    workbook = Workbook(write_only=True)
    sheet = workbook.create_sheet("Dados")
    for row in range(rows):
        values = [f"v{row}-{column}" for column in range(columns)]
        for index in range(min(formulas, columns)):
            values[index] = "=1+1"
        sheet.append(values)
        formulas = max(0, formulas - columns)
    for name, content in (extra or {}).items():
        other = workbook.create_sheet(name)
        for line in content:
            other.append(line)
    stream = BytesIO()
    workbook.save(stream)
    return stream.getvalue()


def _with_sheet_xml(content: bytes, sheet_xml: bytes) -> bytes:
    target = BytesIO()
    with (
        zipfile.ZipFile(BytesIO(content)) as original,
        zipfile.ZipFile(target, "w", zipfile.ZIP_DEFLATED) as modified,
    ):
        for entry in original.infolist():
            payload = original.read(entry)
            if entry.filename == "xl/worksheets/sheet1.xml":
                payload = sheet_xml
            modified.writestr(entry, payload)
    return target.getvalue()


def _sheet_xml(dimension: str | None, rows: str) -> bytes:
    declared = f'<dimension ref="{dimension}"/>' if dimension else ""
    return (
        f'<worksheet xmlns="{SHEET_NS}">{declared}'
        f"<sheetData>{rows}</sheetData></worksheet>"
    ).encode()


def _inspect(tmp_path, content: bytes, source: str = "N-FP"):
    upload = UploadFile(
        file=BytesIO(content),
        filename="book.xlsx",
        headers=Headers({"content-type": XLSX_MIME}),
    )
    return asyncio.run(
        receive_and_inspect(
            upload, source=source, execution_id="xlsx-admission", root=tmp_path
        )
    )


@pytest.fixture(autouse=True)
def _clean_budget(monkeypatch):
    for name in (
        "INGESTION_MAX_SHEETS",
        "INGESTION_MAX_ROWS",
        "INGESTION_MAX_COLUMNS",
        "INGESTION_MAX_CELLS",
        "INGESTION_MAX_FORMULAS",
        "INGESTION_MAX_SECONDS",
    ):
        monkeypatch.delenv(name, raising=False)


@pytest.mark.parametrize("value", [None, ""])
def test_rejects_every_workbook_when_no_sheet_is_authorized(
    tmp_path, monkeypatch, value
) -> None:
    if value is None:
        monkeypatch.delenv("INGESTION_ALLOWED_SHEETS", raising=False)
    else:
        monkeypatch.setenv("INGESTION_ALLOWED_SHEETS", value)

    result = _inspect(tmp_path, _workbook_bytes())

    assert result.reason_code == "no_authorized_sheet"
    assert "INGESTION_ALLOWED_SHEETS" in result.reason_details["guidance"]


def test_memory_estimate_bounds_narrow_sheets_with_many_rows(
    tmp_path, monkeypatch
) -> None:
    per_row = ROW_COST_BYTES + 2 * CELL_COST_BYTES
    monkeypatch.setenv("INGESTION_MAX_MEMORY_BYTES", str(per_row * 50))

    at_limit = _inspect(tmp_path, _workbook_bytes(rows=50, columns=2))
    above_limit = _inspect(tmp_path, _workbook_bytes(rows=51, columns=2))

    assert at_limit.accepted, at_limit.reason_code
    assert above_limit.reason_code == "workbook_memory_limit"
    assert above_limit.reason_details["limit"] == per_row * 50


def test_memory_check_does_not_reject_a_large_sheet_by_its_prefix(
    tmp_path, monkeypatch
) -> None:
    # Past the rule-detail deferral threshold each row is cheaper, so a prefix of
    # the sheet must not be priced with the pre-deferral row cost.
    monkeypatch.setattr(ingestion_budget, "RULE_DETAIL_DEFER_ROWS", 40)
    limit = 100 * DEFERRED_ROW_COST_BYTES + 200 * CELL_COST_BYTES
    monkeypatch.setenv("INGESTION_MAX_MEMORY_BYTES", str(limit))

    result = _inspect(tmp_path, _workbook_bytes(rows=100, columns=2))

    assert result.accepted, result.reason_code


def test_accepts_a_workbook_within_every_limit(tmp_path) -> None:
    result = _inspect(tmp_path, _workbook_bytes(rows=5, columns=4))

    assert result.accepted
    assert result.reason_details is None


@pytest.mark.parametrize(
    ("variable", "limit", "build", "reason"),
    [
        (
            "INGESTION_MAX_ROWS",
            7,
            lambda n: _workbook_bytes(rows=n),
            "workbook_row_limit",
        ),
        (
            "INGESTION_MAX_COLUMNS",
            6,
            lambda n: _workbook_bytes(columns=n),
            "workbook_column_limit",
        ),
        (
            "INGESTION_MAX_FORMULAS",
            5,
            lambda n: _workbook_bytes(rows=4, columns=4, formulas=n),
            "workbook_formula_limit",
        ),
        (
            "INGESTION_MAX_CELLS",
            24,
            lambda n: _workbook_bytes(rows=n // 4, columns=4),
            "workbook_cell_limit",
        ),
    ],
    ids=("rows", "columns", "formulas", "cells"),
)
def test_limit_boundary_accepts_the_limit_and_rejects_one_more(
    tmp_path, monkeypatch, variable, limit, build, reason
) -> None:
    monkeypatch.setenv(variable, str(limit))

    at_limit = _inspect(tmp_path, build(limit))
    above_limit = _inspect(tmp_path, build(limit + (4 if "CELLS" in variable else 1)))

    assert at_limit.accepted, at_limit.reason_code
    assert not above_limit.accepted
    assert above_limit.reason_code == reason
    assert above_limit.reason_details["limit"] == limit
    assert above_limit.reason_details["guidance"]


def test_rejects_more_sheets_than_the_budget(tmp_path, monkeypatch) -> None:
    monkeypatch.setenv("INGESTION_MAX_SHEETS", "2")
    content = _workbook_bytes(extra={"B": [["x"]], "C": [["y"]]})

    result = _inspect(tmp_path, content)

    assert result.reason_code == "workbook_sheet_limit"
    assert result.reason_details["observed"] == 3


def test_rejects_an_abusive_declared_dimension_before_iterating(tmp_path) -> None:
    content = _with_sheet_xml(
        _workbook_bytes(),
        _sheet_xml("A1:XFD1048576", '<row r="1"><c r="A1" t="inlineStr"/></row>'),
    )

    result = _inspect(tmp_path, content)

    assert result.reason_code == "workbook_dimension_limit"
    assert result.reason_details["observed"] == 16384 * 1048576


def test_counts_cells_when_the_sheet_declares_no_dimension(
    tmp_path, monkeypatch
) -> None:
    monkeypatch.setenv("INGESTION_MAX_CELLS", "3")
    cells = "".join(f'<c r="{column}1" t="inlineStr"/>' for column in "ABCD")
    content = _with_sheet_xml(
        _workbook_bytes(), _sheet_xml(None, f"<row>{cells}</row>")
    )

    result = _inspect(tmp_path, content)

    assert result.reason_code == "workbook_cell_limit"
    assert result.reason_details["observed"] == 4


def test_budget_only_applies_to_authorized_sheets(tmp_path, monkeypatch) -> None:
    monkeypatch.setenv("INGESTION_ALLOWED_SHEETS", "Pequena")
    monkeypatch.setenv("INGESTION_MAX_ROWS", "3")
    content = _workbook_bytes(rows=50, extra={"Pequena": [["a", "b"], ["c", "d"]]})

    result = _inspect(tmp_path, content)

    assert result.accepted, result.reason_code


def test_rejects_when_no_sheet_is_authorized(tmp_path, monkeypatch) -> None:
    monkeypatch.setenv("INGESTION_ALLOWED_SHEETS", "Inexistente")

    result = _inspect(tmp_path, _workbook_bytes())

    assert result.reason_code == "no_authorized_sheet"


def test_dangerous_formula_is_blocked_even_in_unauthorized_sheets(
    tmp_path, monkeypatch
) -> None:
    monkeypatch.setenv("INGESTION_ALLOWED_SHEETS", "Dados")
    content = _workbook_bytes(extra={"Oculta": [['=WEBSERVICE("https://invalid")']]})

    result = _inspect(tmp_path, content)

    assert result.reason_code == "dangerous_formula"


def test_time_budget_stops_the_inspection(tmp_path, monkeypatch) -> None:
    monkeypatch.setenv("INGESTION_MAX_SECONDS", "0.000001")

    result = _inspect(tmp_path, _workbook_bytes(rows=2000, columns=10))

    assert result.reason_code == "ingestion_time_limit"


def test_rejects_corrupted_sheet_xml_without_raising(tmp_path) -> None:
    content = _with_sheet_xml(
        _workbook_bytes(), _sheet_xml("A1:A1", '<row r="1"><c r="A1"></row>')
    )

    result = _inspect(tmp_path, content)

    assert result.reason_code == "corrupted_file"


def test_rejects_entity_declarations_in_sheet_xml(tmp_path) -> None:
    sheet = b'<!DOCTYPE x [<!ENTITY a "aaaa">]>' + _sheet_xml(
        "A1:A1", '<row r="1"><c r="A1"/></row>'
    )

    result = _inspect(tmp_path, _with_sheet_xml(_workbook_bytes(), sheet))

    assert result.reason_code == "corrupted_file"


def test_marker_split_across_read_blocks_is_still_detected() -> None:
    marker = b'TargetMode="External"'
    payload = b"a" * (READ_BLOCK - 5) + marker + b"b" * 10
    reader = ScanningReader(
        BytesIO(payload),
        IngestionBudget().start(),
        ByteLedger(limit=10 * READ_BLOCK),
    )

    with pytest.raises(PackageViolation) as excinfo:
        reader.drain()

    assert excinfo.value.code == "external_link"


def test_decompressed_total_is_enforced_while_streaming() -> None:
    reader = ScanningReader(
        BytesIO(b"a" * (3 * READ_BLOCK)),
        IngestionBudget().start(),
        ByteLedger(limit=2 * READ_BLOCK),
    )

    with pytest.raises(PackageViolation) as excinfo:
        reader.drain()

    assert excinfo.value.code == "archive_uncompressed_limit"


def test_inspection_memory_does_not_grow_with_the_sheet_size(tmp_path) -> None:
    content = _workbook_bytes(rows=40_000, columns=8)
    with zipfile.ZipFile(BytesIO(content)) as archive:
        uncompressed = sum(entry.file_size for entry in archive.infolist())
    assert uncompressed > 8 * 1024 * 1024

    tracemalloc.start()
    try:
        result = _inspect(tmp_path, content)
        _, peak = tracemalloc.get_traced_memory()
    finally:
        tracemalloc.stop()

    assert result.accepted, result.reason_code
    assert peak < uncompressed / 2, (peak, uncompressed)
