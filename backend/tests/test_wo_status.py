from __future__ import annotations

import json
import sys
from collections import Counter
from io import BytesIO
from pathlib import Path

import pytest
from openpyxl import load_workbook
from openpyxl.utils.datetime import from_excel

from app.execution import validate_transition
from app.pipeline import read_source, run_pipeline_batch

ROOT = Path(__file__).resolve().parents[2]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

from scripts.generate_wo_status_fixture import SENTINEL, build_workbook  # noqa: E402


class Repository:
    def __init__(self) -> None:
        self.state = "pending"

    def commit_pipeline(self, execution_id: str, result: dict) -> None:
        self.transition_execution(execution_id, result["status"], result["status"])

    def transition_execution(self, execution_id: str, target: str, reason: str) -> None:
        validate_transition(self.state, target)
        self.state = target


def _edit(content: bytes, changes: dict[tuple[str, str], object]) -> bytes:
    book = load_workbook(BytesIO(content))
    for (sheet, cell), value in changes.items():
        book[sheet][cell] = value
    output = BytesIO()
    book.save(output)
    return output.getvalue()


def _process(tmp_path, content: bytes):
    path = tmp_path / "wo-status.xlsx"
    path.write_bytes(content)
    tables, read_issues = read_source(path, ".xlsx", "WO Status")
    result = run_pipeline_batch(
        execution_id="exec-wo-status",
        inputs=[
            {
                "file_name": path.name,
                "source": "WO Status",
                "source_file_id": 7,
                "tables": tables,
                "read_issues": read_issues,
            }
        ],
        repository=Repository(),
        classified_at="2026-08-28T12:00:00+00:00",
    )
    return tables, result


def _issues(result: dict, code: str, sheet: str | None = None) -> list[dict]:
    return [
        issue
        for issue in result["issues"]
        if issue["code"] == code and (sheet is None or issue["sheet"] == sheet)
    ]


def _records(result: dict, sheet: str) -> list[dict]:
    return [r for r in result["normalized_records"] if r["sheet"] == sheet]


def _stats(tables) -> dict[str, Counter]:
    collected = {}
    for sheet, _, rows in tables:
        list(rows)
        collected[sheet] = rows.stats
    return collected


@pytest.fixture(scope="module")
def clean():
    return build_workbook()


@pytest.fixture(scope="module")
def edge():
    return build_workbook(edge=True)


def test_only_the_contract_sheets_are_ingested_and_the_rest_is_reported(
    tmp_path, clean
) -> None:
    content, _ = clean

    tables, result = _process(tmp_path, content)

    assert [sheet for sheet, _, _ in tables] == [
        "NFP",
        "OWM SN",
        "LOT PENDING_OQC",
        "CTNR STATUS",
    ]
    ignored = {issue["sheet"] for issue in _issues(result, "sheet_ignored")}
    assert len(ignored) == 13
    assert {"Mercosul", "Analysis QTY", "Sales", "Master"} <= ignored
    everything = json.dumps(
        [result["normalized_records"], result["imported_records"]], default=str
    )
    assert SENTINEL not in everything


def test_each_sheet_keeps_its_logical_source_and_provenance(tmp_path, clean) -> None:
    content, _ = clean

    _, result = _process(tmp_path, content)

    by_sheet = {r["sheet"]: r["source"] for r in result["normalized_records"]}
    assert by_sheet == {"NFP": "N-FP", "OWM SN": "OWM", "CTNR STATUS": "TMS"}
    record = _records(result, "NFP")[0]
    assert record["source_file_id"] == 7
    assert record["row"] == 4
    transformation = record["transformations"][0]
    assert transformation["adapter"] == "wo_status@1.0.0"
    assert transformation["source_column"] == "U:Plant"
    assert set(result["sources"]) == {"N-FP", "OWM", "GMES/OQC", "TMS"}


def test_mirrored_headers_resolve_to_the_literal_columns(tmp_path, clean) -> None:
    content, _ = clean

    _, result = _process(tmp_path, content)

    assert _issues(result, "formula_not_allowed") == []
    first = _records(result, "NFP")[0]["values"]
    assert first["demand_id"] == "D000000"
    assert first["plant"] == "SYN"
    assert first["lot_quantity"] == 10


def test_identifiers_keep_zeros_and_numbers_are_flagged(tmp_path, edge) -> None:
    content, _ = edge

    _, result = _process(tmp_path, content)

    values = [r["values"] for r in _records(result, "NFP")]
    assert values[0]["workorder_number"] == "0000000"
    by_demand = {v["demand_id"]: v for v in values}
    assert by_demand["D000006"]["workorder_number"] == "0000006"
    assert by_demand["D000008"]["workorder_number"] is None
    numeric = next(v for v in values if v["demand_id"] == "1234567")
    assert numeric["lot_quantity"] == 5
    assert len(_issues(result, "identifier_stored_as_number", "NFP")) == 1


def test_formula_negative_and_missing_identity_reject_only_their_rows(
    tmp_path, edge
) -> None:
    content, expected = edge

    _, result = _process(tmp_path, content)

    nfp = expected["NFP"]
    rejected = {
        issue["row"]
        for code in ("formula_not_allowed", "invalid_quantity", "missing_identity")
        for issue in _issues(result, code, "NFP")
    }
    assert len(rejected) == nfp["rejected_adapter"] == 3
    emitted_rows = {r["row"] for r in _records(result, "NFP")}
    assert emitted_rows.isdisjoint(rejected)
    assert len(emitted_rows) == nfp["emitted"] - nfp["rejected_adapter"]


def test_quantities_dates_and_unknown_states_stay_explicit(tmp_path, edge) -> None:
    content, _ = edge

    _, result = _process(tmp_path, content)

    ctnr = {r["values"]["container_number"]: r for r in _records(result, "CTNR STATUS")}
    serial = ctnr["CTNR-EDGE-SERIAL"]["values"]
    assert serial["load_plan_date"] == from_excel(46000).date().isoformat()
    assert serial["loading_date"] == "2026-12-25"
    assert "CTNR-EDGE-QTY" not in ctnr
    warning = _issues(result, "unparsed_date", "CTNR STATUS")[0]
    assert warning["severity"] == "warning"
    original = ctnr["CTNR-EDGE-SERIAL"]["original_values"]["load_confirmed_at"]
    assert original == "amanhã"
    unknown = _issues(result, "unknown_state", "NFP")
    assert len(unknown) == 1 and unknown[0]["original_value"] == (
        "Totalmente-Desconhecido"
    )


def test_trailing_formula_rows_are_ignored_and_gaps_are_rejected(
    tmp_path, edge
) -> None:
    content, expected = edge

    tables, result = _process(tmp_path, content)

    stats = _stats(tables)["OWM SN"]
    assert stats["ignored_tail"] == expected["OWM SN"]["ignored_tail"] == 8
    assert stats["rejected_gap"] == 1
    assert len(_issues(result, "gap_in_active_region", "OWM SN")) == 1
    serials = [r["values"]["serial_number"] for r in _records(result, "OWM SN")]
    assert "SN-EDGE-A" in serials
    assert [i["column"] for i in _issues(result, "duplicate_serial", "OWM SN")] == ["O"]


def test_blocked_sheet_rejects_every_substantive_row_and_ingests_nothing(
    tmp_path, clean
) -> None:
    content, expected = clean

    tables, result = _process(tmp_path, content)

    blocked = _issues(result, "contract_blocked", "LOT PENDING_OQC")
    assert len(blocked) == expected["LOT PENDING_OQC"]["rejected_blocked"] == 6
    assert _records(result, "LOT PENDING_OQC") == []
    assert "GMES/OQC" not in {r["source"] for r in result["normalized_records"]}
    stats = _stats(tables)["LOT PENDING_OQC"]
    assert stats["ignored_non_substantive"] == 10


def test_counts_match_what_the_fixture_was_built_with(tmp_path, edge) -> None:
    content, expected = edge

    tables, _ = _process(tmp_path, content)

    stats = _stats(tables)
    for sheet, counts in expected.items():
        for name, value in counts.items():
            assert stats[sheet][name] == value, (sheet, name, stats[sheet])


def test_header_drift_blocks_only_the_drifted_sheet(tmp_path, clean) -> None:
    content, _ = clean
    drifted = _edit(content, {("NFP", "Y3"): "Demand Code"})

    _, result = _process(tmp_path, drifted)

    drift = _issues(result, "contract_drift")
    assert [(i["sheet"], i["column"], i["row"]) for i in drift] == [("NFP", "Y", 3)]
    assert _records(result, "NFP") == []
    assert len(_records(result, "OWM SN")) > 0
    assert result["blocking"] is True


def test_authorized_sheets_follow_the_environment_when_configured(
    tmp_path, clean, monkeypatch
) -> None:
    content, _ = clean
    monkeypatch.setenv("INGESTION_ALLOWED_SHEETS_WO_STATUS", "NFP|CTNR STATUS")

    tables, result = _process(tmp_path, content)

    assert [sheet for sheet, _, _ in tables] == ["NFP", "CTNR STATUS"]
    not_authorized = {i["sheet"] for i in _issues(result, "sheet_not_authorized")}
    assert not_authorized == {"OWM SN", "LOT PENDING_OQC"}


def test_missing_authorized_sheet_is_reported(tmp_path, clean) -> None:
    content, _ = clean
    book = load_workbook(BytesIO(content))
    del book["CTNR STATUS"]
    output = BytesIO()
    book.save(output)

    _, result = _process(tmp_path, output.getvalue())

    assert [i["sheet"] for i in _issues(result, "sheet_missing")] == ["CTNR STATUS"]


def test_reprocessing_is_deterministic(tmp_path, clean) -> None:
    content, _ = clean

    _, first = _process(tmp_path, content)
    _, second = _process(tmp_path, content)

    def digest(result: dict) -> str:
        return json.dumps(
            [result["normalized_records"], result["issues"], result["summary"]],
            default=str,
            sort_keys=True,
        )

    assert digest(first) == digest(second)


def test_csv_is_not_accepted_for_the_composite_source(tmp_path) -> None:
    path = tmp_path / "wo.csv"
    path.write_text("workorder\nWO-1\n", encoding="utf-8")

    with pytest.raises(ValueError, match="XLSX"):
        read_source(path, ".csv", "WO Status")
