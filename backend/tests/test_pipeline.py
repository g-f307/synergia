import tracemalloc
from pathlib import Path

import pytest
from openpyxl import Workbook

from app.execution import validate_transition
from app.pipeline import read_source, run_pipeline, run_pipeline_batch


class RecordingRepository:
    def __init__(self) -> None:
        self.commits: list[tuple[str, dict]] = []
        self.state = "pending"
        self.transitions: list[str] = []

    def commit_pipeline(self, execution_id: str, result: dict) -> None:
        self.commits.append((execution_id, result))
        self.transition_execution(execution_id, result["status"], result["status"])

    def transition_execution(self, execution_id: str, target: str, reason: str) -> None:
        validate_transition(self.state, target)
        self.state = target
        self.transitions.append(target)


def run_csv(
    tmp_path: Path,
    content: str,
    *,
    source: str = "N-FP",
    organizations: set[str] | None = None,
) -> tuple[dict, RecordingRepository]:
    path = tmp_path / "input.csv"
    path.write_text(content, encoding="utf-8")
    repository = RecordingRepository()
    tables, read_issues = read_source(path, ".csv", source)
    result = run_pipeline(
        execution_id="exec-pipeline",
        file_name="input.csv",
        source=source,
        repository=repository,
        tables=tables,
        read_issues=read_issues,
        source_file_id=1,
        classified_at="2026-08-28T12:00:00+00:00",
        known_organizations=organizations,
    )
    return result, repository


def test_valid_file_runs_all_stages_and_preserves_traceability(tmp_path) -> None:
    result, repository = run_csv(
        tmp_path,
        "workorder,serial,planned_date,planned_quantity\n"
        "0000123,0000456,27/08/2026,10\n",
    )

    assert result["status"] == "completed"
    assert result["summary"] == {
        "rows_read": 1,
        "valid_records": 1,
        "rejected_records": 0,
        "normalized_records": 1,
        "errors": 0,
        "warnings": 0,
    }
    imported = result["imported_records"][0]
    normalized = result["normalized_records"][0]
    assert (imported["sheet"], imported["row"]) == (
        normalized["sheet"],
        normalized["row"],
    )
    assert normalized["values"]["workorder_number"] == "0000123"
    assert normalized["values"]["serial_number"] == "0000456"
    assert normalized["original_values"]["planned_date"] == "27/08/2026"
    assert normalized["values"]["planned_date"] == "2026-08-27"
    assert normalized["execution_id"] == "exec-pipeline"
    assert normalized["source_file_id"] == 1
    assert (
        result["processing"]["consolidation"]["workorders"][0]["provenance"][
            "planned_quantity"
        ][0]["source_file_id"]
        == 1
    )
    assert result["processing"]["summary"] == {
        "eligible_normalized_records": 1,
        "consolidated_workorders": 1,
        "consolidated_lots": 0,
        "consolidated_serials": 1,
        "consolidated_organizations": 0,
        "consolidation_issues": 1,
        "failed_workorders": 0,
        "classifications": 0,
        "active_pending_items": 0,
        "classifications_by_rule": {},
        "consolidated_quantities": {
            "planned_quantity": {"known_workorders": 1, "total": 10},
            "produced_quantity": {"known_workorders": 0, "total": None},
            "received_quantity": {"known_workorders": 0, "total": None},
            "released_quantity": {"known_workorders": 0, "total": None},
            "pending_quantity": {"known_workorders": 0, "total": None},
            "retained_quantity": {"known_workorders": 0, "total": None},
        },
    }
    assert len(repository.commits) == 1
    assert repository.transitions == [
        "validating",
        "normalizing",
        "consolidating",
        "applying_rules",
        "completed",
    ]


def test_record_error_rejects_only_affected_line(tmp_path) -> None:
    result, _ = run_csv(
        tmp_path,
        "workorder,planned_quantity,planned_date\n"
        "WO-OK,10,2026-08-27\n"
        "WO-BAD,invalid,not-a-date\n"
        "WO-ALSO-OK,2,28/08/2026\n",
    )

    assert result["status"] == "completed_with_errors"
    assert result["blocking"] is False
    assert result["summary"]["valid_records"] == 2
    assert result["summary"]["rejected_records"] == 1
    assert [record["row"] for record in result["normalized_records"]] == [2, 4]
    assert {
        item["workorder_number"]
        for item in result["processing"]["consolidation"]["workorders"]
    } == {"WO-OK", "WO-ALSO-OK"}
    assert {issue["scope"] for issue in result["issues"]} == {"record"}


@pytest.mark.parametrize(
    "content",
    ["status\nopen\n", "workorder\n", ""],
)
def test_structural_error_blocks_the_file(tmp_path, content) -> None:
    result, repository = run_csv(tmp_path, content)

    assert result["status"] == "validation_failed"
    assert result["blocking"] is True
    assert result["summary"]["normalized_records"] == 0
    assert any(issue["scope"] == "structure" for issue in result["issues"])
    assert repository.transitions == ["validating", "validation_failed"]


def test_required_field_and_unknown_organization_are_row_errors(tmp_path) -> None:
    result, _ = run_csv(
        tmp_path,
        "workorder,organization\n,KNOWN\nWO-2,UNKNOWN\nWO-3,KNOWN\n",
        organizations={"KNOWN"},
    )

    assert result["summary"]["rejected_records"] == 2
    assert result["summary"]["normalized_records"] == 1
    assert {issue["code"] for issue in result["issues"]} == {
        "required_field",
        "unknown_organization",
    }


def test_unknown_state_and_oqc_flag_are_non_blocking_warnings(tmp_path) -> None:
    result, _ = run_csv(
        tmp_path,
        "workorder,status,oqc_flag\nWO-1,new-state,perhaps\n",
        source="GMES/OQC",
    )

    assert result["summary"]["normalized_records"] == 1
    assert result["summary"]["warnings"] == 2
    assert all(issue["severity"] == "warning" for issue in result["issues"])


def test_unexpected_normalization_failure_does_not_commit(
    tmp_path, monkeypatch
) -> None:
    path = tmp_path / "input.csv"
    path.write_text("workorder\nWO-1\n", encoding="utf-8")
    repository = RecordingRepository()
    tables, read_issues = read_source(path, ".csv", "N-FP")

    def fail_normalization(*_args, **_kwargs):
        raise RuntimeError("synthetic failure")

    monkeypatch.setattr("app.pipeline.normalize_tables", fail_normalization)
    with pytest.raises(RuntimeError, match="synthetic failure"):
        run_pipeline(
            execution_id="exec-failure",
            file_name="input.csv",
            source="N-FP",
            repository=repository,
            tables=tables,
            read_issues=read_issues,
            source_file_id=1,
            classified_at="2026-08-28T12:00:00+00:00",
        )

    assert repository.commits == []


def test_artifact_failure_does_not_commit_pipeline(tmp_path) -> None:
    path = tmp_path / "input.csv"
    path.write_text("workorder\nWO-1\n", encoding="utf-8")
    tables, read_issues = read_source(path, ".csv", "N-FP")
    repository = RecordingRepository()

    def fail_artifacts(_result: dict) -> None:
        raise OSError("synthetic artifact failure")

    with pytest.raises(OSError, match="synthetic artifact failure"):
        run_pipeline(
            execution_id="exec-artifact-failure",
            file_name="input.csv",
            source="N-FP",
            repository=repository,
            tables=tables,
            read_issues=read_issues,
            source_file_id=1,
            classified_at="2026-08-28T12:00:00+00:00",
            prepare_commit=fail_artifacts,
        )

    assert repository.commits == []


def test_reads_xlsx_from_temporary_path_without_xlsx_suffix(tmp_path) -> None:
    path = tmp_path / "upload.upload"
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.append(["workorder", "status"])
    worksheet.append(["WO-1", "open"])
    workbook.save(path)
    workbook.close()

    tables, read_issues = read_source(path, ".xlsx", "N-FP")

    assert read_issues == []
    assert [(sheet, headers, list(rows)) for sheet, headers, rows in tables] == [
        ("Sheet", ["workorder", "status"], [["WO-1", "open"]])
    ]
    assert not isinstance(tables[0][2], list)


def test_reads_multisheet_quality_workbook_with_titles_and_real_row_numbers(
    tmp_path,
) -> None:
    path = tmp_path / "quality.upload"
    workbook = Workbook()
    first = workbook.active
    first.title = "Inspection_A"
    first.append(["INSPECTION REPORT"])
    first.append([])
    first.append(["lote_id", "produto", "status", "data"])
    first.append(["0000007", "MODEL-A", "approved", "01/09/2026"])
    second = workbook.create_sheet("Inspection_B")
    second.append(["ANOTHER REPORT"])
    second.append([])
    second.append(["lote_id", "produto", "status", "data"])
    second.append(["0000008", "MODEL-B", "hold", "02/09/2026"])
    workbook.save(path)
    workbook.close()

    tables, read_issues = read_source(path, ".xlsx", "GMES/OQC")

    assert read_issues == []
    assert [table[0] for table in tables] == ["Inspection_A", "Inspection_B"]
    assert [table[2][0].row_number for table in tables] == [4, 4]

    result = run_pipeline_batch(
        execution_id="exec-reference-contract",
        inputs=[
            {
                "file_name": "plan.csv",
                "source": "N-FP",
                "source_file_id": 1,
                "tables": [
                    (
                        "CSV",
                        ["workorder_number", "lot_number", "planned_quantity"],
                        [["WO-7", "0000007", 1], ["WO-8", "0000008", 1]],
                    )
                ],
                "read_issues": [],
            },
            {
                "file_name": "quality.xlsx",
                "source": "GMES/OQC",
                "source_file_id": 2,
                "tables": tables,
                "read_issues": read_issues,
            },
        ],
        repository=RecordingRepository(),
        classified_at="2026-09-01T12:00:00+00:00",
    )

    assert result["status"] == "completed"
    quality_rows = [
        record["row"]
        for record in result["normalized_records"]
        if record["source"] == "GMES/OQC"
    ]
    assert quality_rows == [4, 4]
    assert {
        item["workorder_number"]
        for item in result["processing"]["consolidation"]["workorders"]
    } == {"WO-7", "WO-8"}
    assert result["normalized_records"][2]["values"]["lot_number"] == "0000007"


def test_batch_pipeline_uses_real_file_ids_for_multisource_provenance() -> None:
    repository = RecordingRepository()
    result = run_pipeline_batch(
        execution_id="exec-batch",
        inputs=[
            {
                "file_name": "plan.csv",
                "source": "N-FP",
                "source_file_id": 41,
                "tables": [("CSV", ["workorder", "planned_quantity"], [["WO", 10]])],
                "read_issues": [],
            },
            {
                "file_name": "receiving.csv",
                "source": "OWM",
                "source_file_id": 42,
                "tables": [("CSV", ["workorder", "planned_quantity"], [["WO", 12]])],
                "read_issues": [],
            },
        ],
        repository=repository,
        classified_at="2026-08-28T12:00:00+00:00",
    )

    workorder = result["processing"]["consolidation"]["workorders"][0]
    assert workorder["planned_quantity"] == 10
    assert {
        origin["source_file_id"]
        for origin in workorder["provenance"]["planned_quantity"]
    } == {41, 42}
    assert (
        result["processing"]["summary"]["classifications_by_rule"]["source_divergence"]
        == 1
    )


def _narrow_workbook(path: Path, rows: int, workorders: int) -> None:
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.append(["workorder", "serial", "status"])
    for index in range(rows):
        worksheet.append([f"WO-{index % workorders:05d}", f"SN-{index:07d}", "open"])
    workbook.save(path)
    workbook.close()


def test_full_pipeline_memory_of_a_narrow_sheet_stays_under_the_budget_estimate(
    tmp_path,
) -> None:
    rows = 4_000
    path = tmp_path / "estreita.xlsx"
    _narrow_workbook(path, rows, workorders=200)
    tables, read_issues = read_source(path, ".xlsx", "OWM")
    budget = tables[0][2].run.budget
    repository = RecordingRepository()

    tracemalloc.start()
    try:
        result = run_pipeline(
            execution_id="exec-narrow",
            file_name="estreita.xlsx",
            source="OWM",
            repository=repository,
            tables=tables,
            read_issues=read_issues,
            source_file_id=1,
            classified_at="2026-08-28T12:00:00+00:00",
        )
        _, peak = tracemalloc.get_traced_memory()
    finally:
        tracemalloc.stop()

    assert result["summary"]["rows_read"] == rows
    assert peak <= budget.estimated_memory(rows + 1, (rows + 1) * 3), peak


def test_consolidation_stage_honours_the_time_budget_in_the_real_code(
    tmp_path, monkeypatch
) -> None:
    from app import ingestion_budget, pipeline
    from app.ingestion_budget import IngestionLimitError

    path = tmp_path / "pequena.xlsx"
    _narrow_workbook(path, rows=30, workorders=3)
    tables, read_issues = read_source(path, ".xlsx", "OWM")
    real_process = pipeline.process_normalized_records

    def expire_then_process(records, **kwargs):
        monkeypatch.setattr(ingestion_budget.time, "monotonic", lambda: 10**9)
        return real_process(records, **kwargs)

    monkeypatch.setattr(pipeline, "process_normalized_records", expire_then_process)

    with pytest.raises(IngestionLimitError) as excinfo:
        run_pipeline(
            execution_id="exec-slow",
            file_name="pequena.xlsx",
            source="OWM",
            repository=RecordingRepository(),
            tables=tables,
            read_issues=read_issues,
            source_file_id=1,
            classified_at="2026-08-28T12:00:00+00:00",
        )

    assert excinfo.value.code == "ingestion_time_limit"


def _normalized(workorders: int) -> list[dict]:
    from app.normalization import normalize_tables

    rows = [[f"WO-{index:05d}", "open"] for index in range(workorders)]
    records = normalize_tables([("S", ["workorder", "status"], rows)], "N-FP")[
        "records"
    ]
    return [{**record, "execution_id": "e", "source_file_id": 1} for record in records]


@pytest.mark.parametrize("defer", [False, True])
def test_every_long_loop_of_the_real_stages_consults_the_clock(defer) -> None:
    from app.processing import process_normalized_records

    calls = []

    class Expired(Exception):
        pass

    def checkpoint() -> None:
        calls.append(1)
        if len(calls) >= 6:
            raise Expired

    with pytest.raises(Expired):
        process_normalized_records(
            _normalized(700),
            execution_id="e",
            classified_at="2026-08-28T12:00:00+00:00",
            defer_rule_details=defer,
            checkpoint=checkpoint,
        )

    assert len(calls) == 6


def test_batch_records_each_file_duration_when_its_phase_ends(
    tmp_path, monkeypatch
) -> None:
    from app import ingestion_budget, pipeline
    from app.observability.metrics import INGESTION_DURATION

    clock = {"now": 1000.0}
    monkeypatch.setattr(ingestion_budget.time, "monotonic", lambda: clock["now"])
    inputs = []
    for position, source in enumerate(("N-FP", "OWM"), start=1):
        path = tmp_path / f"arquivo-{position}.xlsx"
        _narrow_workbook(path, rows=5, workorders=5)
        tables, read_issues = read_source(path, ".xlsx", source)
        inputs.append(
            {
                "file_name": path.name,
                "source": source,
                "source_file_id": position,
                "tables": tables,
                "read_issues": read_issues,
            }
        )
    real_normalize = pipeline.normalize_tables

    def slow_second_file(tables, source, *args, **kwargs):
        if source == "OWM":
            clock["now"] += 100.0
        return real_normalize(tables, source, *args, **kwargs)

    monkeypatch.setattr(pipeline, "normalize_tables", slow_second_file)
    duration = INGESTION_DURATION.labels("completed")
    before = duration._sum.get()

    run_pipeline_batch(
        execution_id="exec-batch",
        inputs=inputs,
        repository=RecordingRepository(),
        classified_at="2026-08-28T12:00:00+00:00",
    )

    # The first file ended at 0 s and the second at 100 s; neither includes the
    # files after it nor the consolidation.
    assert duration._sum.get() - before == pytest.approx(100.0)


def test_unexpected_failure_is_recorded_with_the_failed_outcome(
    tmp_path, monkeypatch
) -> None:
    from app import pipeline
    from app.observability.metrics import INGESTION_DURATION

    path = tmp_path / "falha.xlsx"
    _narrow_workbook(path, rows=5, workorders=5)
    tables, read_issues = read_source(path, ".xlsx", "OWM")

    def broken(*args, **kwargs):
        raise RuntimeError("falha inesperada")

    monkeypatch.setattr(pipeline, "process_normalized_records", broken)
    failed = INGESTION_DURATION.labels("failed")
    before = sum(bucket.get() for bucket in failed._buckets)

    with pytest.raises(RuntimeError):
        run_pipeline(
            execution_id="exec-broken",
            file_name=path.name,
            source="OWM",
            repository=RecordingRepository(),
            tables=tables,
            read_issues=read_issues,
            source_file_id=1,
            classified_at="2026-08-28T12:00:00+00:00",
        )

    assert sum(bucket.get() for bucket in failed._buckets) == before + 1


def _threshold_case(tmp_path, monkeypatch, data_rows: int, threshold: int = 40):
    from app import ingestion_budget

    monkeypatch.setattr(ingestion_budget, "RULE_DETAIL_DEFER_ROWS", threshold)
    monkeypatch.setattr("app.pipeline.RULE_DETAIL_DEFER_THRESHOLD", threshold)
    # Physical rows include the header, so the streaming lower bound sees one
    # row more than the normalized records.
    cheapest = (
        (threshold + 1) * ingestion_budget.DEFERRED_ROW_COST_BYTES
        + threshold * 2 * ingestion_budget.CELL_COST_BYTES
    )
    monkeypatch.setenv("INGESTION_MAX_MEMORY_BYTES", str(cheapest + 1))
    path = tmp_path / f"limiar-{data_rows}.xlsx"
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.append(["workorder", "status"])
    for index in range(data_rows):
        worksheet.append([f"WO-{index:05d}", "open"])
    workbook.save(path)
    workbook.close()
    tables, read_issues = read_source(path, ".xlsx", "N-FP")
    return run_pipeline(
        execution_id="exec-threshold",
        file_name=path.name,
        source="N-FP",
        repository=RecordingRepository(),
        tables=tables,
        read_issues=read_issues,
        source_file_id=1,
        classified_at="2026-08-28T12:00:00+00:00",
    )


def test_estimate_uses_the_normalized_count_that_decides_deferral(
    tmp_path, monkeypatch
) -> None:
    from app.ingestion_budget import IngestionLimitError

    # 40 physical rows (header included) leave 39 normalized records, which is
    # below the deferral threshold: the expensive regime applies even though the
    # physical row count alone looks like the cheap one.
    with pytest.raises(IngestionLimitError) as excinfo:
        _threshold_case(tmp_path, monkeypatch, data_rows=39)

    assert excinfo.value.code == "workbook_memory_limit"
    assert excinfo.value.observed > excinfo.value.limit


def test_estimate_switches_to_the_cheap_regime_at_the_threshold(
    tmp_path, monkeypatch
) -> None:
    result = _threshold_case(tmp_path, monkeypatch, data_rows=40)

    assert result["summary"]["rows_read"] == 40
    assert result["processing"]["classifications"]["rule_details_deferred"] is True
