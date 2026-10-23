from __future__ import annotations

import json
import subprocess
import sys
from io import BytesIO
from pathlib import Path

from openpyxl import load_workbook

from app.wo_status_report import build_report, compare_persisted

ROOT = Path(__file__).resolve().parents[2]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

from scripts.generate_wo_status_fixture import SENTINEL, build_workbook  # noqa: E402

SCRIPT = ROOT / "scripts" / "validate_wo_status.py"


def _write(tmp_path, content: bytes) -> Path:
    path = tmp_path / "wo-status.xlsx"
    path.write_bytes(content)
    return path


def _by_sheet(report: dict) -> dict[str, dict]:
    return {sheet["sheet"]: sheet for sheet in report["sheets"]}


def test_report_reconciles_every_authorized_sheet(tmp_path) -> None:
    content, expected = build_workbook(edge=True)

    report = build_report(_write(tmp_path, content))

    assert report["balanced"] is True
    assert report["adapter"] == "wo_status@1.0.0"
    assert report["contains_source_values"] is False
    sheets = _by_sheet(report)
    assert sheets["NFP"]["read"] == expected["NFP"]["scanned"]
    assert sheets["NFP"]["rejected"] == expected["NFP"]["rejected_adapter"]
    assert sheets["NFP"]["ignored_total"] == expected["NFP"]["ignored_blank"]
    assert sheets["OWM SN"]["ignored"]["ignored_tail"] == 8
    assert sheets["OWM SN"]["duplicated"] == 1
    assert sheets["LOT PENDING_OQC"]["status"] == "blocked"
    assert sheets["LOT PENDING_OQC"]["normalized"] == 0
    assert sheets["LOT PENDING_OQC"]["rejected"] == 6
    assert sheets["CTNR STATUS"]["rejected"] == 1
    for sheet in sheets.values():
        assert sheet["read"] == sheet["ignored_total"] + sheet["entered_validation"]
        assert sheet["entered_validation"] == sheet["normalized"] + sheet["rejected"]
        assert sheet["imported"] == sheet["entered_validation"]


def test_report_lists_ignored_sheets_and_never_prints_values(tmp_path) -> None:
    content, _ = build_workbook(edge=True)

    report = build_report(_write(tmp_path, content))

    assert len(report["ignored_sheets"]) == 13
    text = json.dumps(report, ensure_ascii=False)
    assert SENTINEL not in text
    for value in ("D000001", "SN00000001", "CTNR0000001", "usuario.sintetico"):
        assert value not in text


def test_report_flags_header_drift_and_missing_sheets(tmp_path) -> None:
    content, _ = build_workbook()
    book = load_workbook(BytesIO(content))
    book["NFP"]["Y3"] = "Outro cabeçalho"
    del book["CTNR STATUS"]
    output = BytesIO()
    book.save(output)

    report = build_report(_write(tmp_path, output.getvalue()))

    sheets = _by_sheet(report)
    assert sheets["NFP"]["status"] == "drift"
    assert sheets["CTNR STATUS"]["status"] == "missing"
    assert sheets["OWM SN"]["status"] == "ingested"
    assert report["blocking"] is True


def test_comparison_detects_missing_extra_and_unexpected_persisted_rows(
    tmp_path,
) -> None:
    content, _ = build_workbook()
    report = build_report(_write(tmp_path, content))
    persisted = {
        sheet["sheet"]: {
            "imported": sheet["imported"],
            "normalized": sheet["normalized"],
        }
        for sheet in report["sheets"]
    }

    assert compare_persisted(report, persisted) == {"matches": True, "differences": []}

    persisted["NFP"]["normalized"] -= 1
    persisted["Sales"] = {"imported": 0, "normalized": 3}
    comparison = compare_persisted(report, persisted)

    assert comparison["matches"] is False
    assert comparison["differences"] == [
        {
            "sheet": "NFP",
            "kind": "normalized",
            "expected": report["sheets"][0]["normalized"],
            "persisted": report["sheets"][0]["normalized"] - 1,
        },
        {
            "sheet": "Sales",
            "kind": "unexpected_sheet",
            "persisted": {"imported": 0, "normalized": 3},
        },
    ]


def test_script_prints_the_report_without_persisting(tmp_path) -> None:
    content, _ = build_workbook()
    path = _write(tmp_path, content)

    completed = subprocess.run(
        [sys.executable, "-I", str(SCRIPT), str(path), "--strict"],
        capture_output=True,
        text=True,
        check=False,
        env={"PATH": "", "PYTHONPATH": ""},
    )

    assert completed.returncode == 0, completed.stderr
    report = json.loads(completed.stdout)
    assert report["balanced"] is True and report["adapter"] == "wo_status@1.0.0"


def test_strict_script_fails_on_contract_drift(tmp_path) -> None:
    content, _ = build_workbook()
    book = load_workbook(BytesIO(content))
    book["OWM SN"]["O2"] = "Serial"
    output = BytesIO()
    book.save(output)
    path = _write(tmp_path, output.getvalue())

    completed = subprocess.run(
        [sys.executable, "-I", str(SCRIPT), str(path), "--strict"],
        capture_output=True,
        text=True,
        check=False,
        env={"PATH": "", "PYTHONPATH": ""},
    )

    assert completed.returncode == 1
    assert json.loads(completed.stdout)["blocking"] is True


def test_script_requires_a_database_to_compare(tmp_path) -> None:
    content, _ = build_workbook()
    path = _write(tmp_path, content)

    completed = subprocess.run(
        [sys.executable, "-I", str(SCRIPT), str(path), "--execution-id", "x"],
        capture_output=True,
        text=True,
        check=False,
        env={"PATH": "", "PYTHONPATH": ""},
    )

    assert completed.returncode == 2
    assert "DATABASE_URL" in completed.stderr
