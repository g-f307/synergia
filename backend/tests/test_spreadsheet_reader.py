from __future__ import annotations

import logging
import tracemalloc

import pytest
from openpyxl import Workbook

from app.ingestion_budget import (
    CELL_COST_BYTES,
    DEFERRED_ROW_COST_BYTES,
    IngestionBudget,
    IngestionLimitError,
)
from app.normalization import normalize_file
from app.spreadsheet_reader import (
    CLOCK_CHECK_INTERVAL,
    HEADER_SCAN_ROWS,
    IngestionRun,
    LazySheetRows,
)
from app.validation import read_tables, validate_tables


def _save(tmp_path, build, name: str = "book.xlsx"):
    workbook = Workbook()
    build(workbook)
    path = tmp_path / name
    workbook.save(path)
    workbook.close()
    return path


def _simple(rows: int, extra: dict[str, int] | None = None):
    def build(workbook: Workbook) -> None:
        sheet = workbook.active
        sheet.title = "Dados"
        sheet.append(["workorder", "status"])
        for index in range(rows):
            sheet.append([f"WO-{index:05d}", "open"])
        for name, size in (extra or {}).items():
            other = workbook.create_sheet(name)
            other.append(["x", "y"])
            for index in range(size):
                other.append([index, index])

    return build


def _table(path, run=None):
    tables, _ = read_tables(path, ".xlsx", "N-FP", run)
    return tables[0]


@pytest.fixture(autouse=True)
def _clean_budget(monkeypatch):
    for name in (
        "INGESTION_MAX_ROWS",
        "INGESTION_MAX_COLUMNS",
        "INGESTION_MAX_CELLS",
        "INGESTION_MAX_ISSUES",
        "INGESTION_PROGRESS_INTERVAL",
    ):
        monkeypatch.delenv(name, raising=False)


def test_xlsx_tables_are_lazy_and_each_pass_rereads_the_file(tmp_path) -> None:
    path = _save(tmp_path, _simple(3))

    tables, issues = read_tables(path, ".xlsx", "N-FP")

    assert issues == []
    ((sheet, headers, rows),) = tables
    assert isinstance(rows, LazySheetRows)
    assert (sheet, headers) == ("Dados", ["workorder", "status"])
    first = [list(row) for row in rows]
    second = [list(row) for row in rows]
    assert first == second == [[f"WO-{i:05d}", "open"] for i in range(3)]
    assert [row.row_number for row in rows] == [2, 3, 4]
    assert rows.count() == 3
    assert rows[0].row_number == 2
    assert rows and rows.header_row == 1


def test_preamble_rows_keep_their_physical_numbers(tmp_path) -> None:
    def build(workbook: Workbook) -> None:
        sheet = workbook.active
        sheet.append(["RELATÓRIO"])
        sheet.append([])
        sheet.append(["workorder", "status"])
        sheet.append(["WO-1", "open"])

    path = _save(tmp_path, build)

    _, _, rows = _table(path)

    assert rows.header_row == 3
    assert [row.row_number for row in rows] == [4]


def test_unauthorized_sheets_are_never_listed_or_read(tmp_path, monkeypatch) -> None:
    monkeypatch.setenv("INGESTION_ALLOWED_SHEETS", "Dados")
    path = _save(tmp_path, _simple(2, {"Auxiliar": 5}))

    tables, _ = read_tables(path, ".xlsx", "N-FP")

    assert [sheet for sheet, _, _ in tables] == ["Dados"]


def test_requires_at_least_one_authorized_sheet(tmp_path, monkeypatch) -> None:
    monkeypatch.setenv("INGESTION_ALLOWED_SHEETS", "Outra")
    path = _save(tmp_path, _simple(2))

    with pytest.raises(IngestionLimitError) as excinfo:
        read_tables(path, ".xlsx", "N-FP")

    assert excinfo.value.code == "no_authorized_sheet"


def test_header_beyond_the_scan_window_is_reported_not_searched_forever(
    tmp_path,
) -> None:
    def build(workbook: Workbook) -> None:
        sheet = workbook.active
        for _ in range(HEADER_SCAN_ROWS + 5):
            sheet.append(["apenas", "texto"])
        sheet.append(["workorder", "status"])

    tables, issues = read_tables(_save(tmp_path, build), ".xlsx", "N-FP")

    assert tables == []
    assert [issue["code"] for issue in issues] == ["invalid_header"]


@pytest.mark.parametrize(
    ("variable", "limit", "rows", "reason"),
    [
        ("INGESTION_MAX_ROWS", 10, 10, "workbook_row_limit"),
        ("INGESTION_MAX_CELLS", 20, 10, "workbook_cell_limit"),
    ],
)
def test_reading_enforces_the_budget_with_boundary_values(
    tmp_path, monkeypatch, variable, limit, rows, reason
) -> None:
    monkeypatch.setenv(variable, str(limit))
    within = _save(tmp_path, _simple(rows - 1), "within.xlsx")
    above = _save(tmp_path, _simple(rows + 5), "above.xlsx")

    _, _, ok_rows = _table(within)
    _, _, big_rows = _table(above)

    assert len(list(ok_rows)) == rows - 1
    with pytest.raises(IngestionLimitError) as excinfo:
        list(big_rows)
    assert excinfo.value.code == reason
    assert excinfo.value.limit == limit


def test_reading_stops_early_when_even_the_cheapest_regime_does_not_fit(
    tmp_path, monkeypatch
) -> None:
    # 19 data rows end at physical row 20 with 38 cells.
    limit = 20 * DEFERRED_ROW_COST_BYTES + 38 * CELL_COST_BYTES
    monkeypatch.setenv("INGESTION_MAX_MEMORY_BYTES", str(limit))
    within = _save(tmp_path, _simple(19), "within.xlsx")
    above = _save(tmp_path, _simple(25), "above.xlsx")

    assert len(list(_table(within)[2])) == 19
    with pytest.raises(IngestionLimitError) as excinfo:
        list(_table(above)[2])
    assert excinfo.value.code == "workbook_memory_limit"


def test_reading_enforces_the_time_budget_every_check_interval(tmp_path) -> None:
    path = _save(tmp_path, _simple(CLOCK_CHECK_INTERVAL + 10))
    run = IngestionRun("N-FP")
    _, _, rows = _table(path, run)
    run.clock = IngestionBudget(max_seconds=0.000001).start()

    with pytest.raises(IngestionLimitError) as excinfo:
        list(rows)

    assert excinfo.value.code == "ingestion_time_limit"


def test_unreadable_workbook_is_reported_as_a_structured_rejection(tmp_path) -> None:
    path = tmp_path / "quebrado.xlsx"
    path.write_bytes(b"PK\x03\x04not a workbook")

    with pytest.raises(IngestionLimitError) as excinfo:
        read_tables(path, ".xlsx", "N-FP")

    assert excinfo.value.code == "unreadable_workbook"


def test_validation_stops_when_the_issue_budget_is_exhausted(tmp_path) -> None:
    def build(workbook: Workbook) -> None:
        sheet = workbook.active
        sheet.append(["workorder", "status"])
        for _ in range(50):
            sheet.append([None, "open"])

    table = _table(_save(tmp_path, build))

    with pytest.raises(IngestionLimitError) as excinfo:
        validate_tables([table], "N-FP", max_issues=10)

    assert excinfo.value.code == "ingestion_issue_limit"
    assert excinfo.value.limit == 10
    assert excinfo.value.observed > 10


def test_duplicate_rows_are_still_detected_without_keeping_row_values() -> None:
    rows = [["WO-1", "open"], ["WO-2", "open"], ["WO-1", " open "]]

    report = validate_tables([("Dados", ["workorder", "status"], rows)], "N-FP")

    duplicates = [i for i in report["issues"] if i["code"] == "duplicate_row"]
    assert [issue["row"] for issue in duplicates] == [4]
    assert "linha 2" in duplicates[0]["reason"]


def test_progress_is_reported_at_the_configured_interval(
    tmp_path, monkeypatch, caplog
) -> None:
    monkeypatch.setenv("INGESTION_PROGRESS_INTERVAL", "100")
    path = _save(tmp_path, _simple(250))
    _, _, rows = _table(path)

    with caplog.at_level(logging.INFO, logger="synergia.telemetry"):
        list(rows)

    events = [r for r in caplog.records if r.getMessage() == "ingestion.progress"]
    assert [r.synergia_fields["processed_count"] for r in events] == [100, 200]
    assert all("file_name" not in r.synergia_fields for r in events)


def test_normalization_reads_xlsx_incrementally(tmp_path) -> None:
    path = _save(tmp_path, _simple(4))

    result = normalize_file(path, ".xlsx", "N-FP")

    assert result["record_count"] == 4
    assert [record["row"] for record in result["records"]] == [2, 3, 4, 5]


def test_validation_and_normalization_memory_does_not_hold_the_sheet(
    tmp_path,
) -> None:
    rows = 12_000

    def build(workbook: Workbook) -> None:
        sheet = workbook.active
        sheet.append(["workorder", "status", "model", "reason", "lot"])
        for index in range(rows):
            sheet.append([f"WO-{index:06d}", "open", "M" * 20, "r" * 40, index])

    path = _save(tmp_path, build)
    table = _table(path)
    _, headers, lazy = table

    tracemalloc.start()
    try:
        materialized = [list(row) for row in lazy]
        _, materialized_peak = tracemalloc.get_traced_memory()
    finally:
        tracemalloc.stop()
    del materialized

    tracemalloc.start()
    try:
        report = validate_tables([table], "N-FP")
        _, lazy_peak = tracemalloc.get_traced_memory()
    finally:
        tracemalloc.stop()

    assert report["row_count"] == rows
    # Validation keeps only a digest per row, never the parsed sheet.
    assert lazy_peak < materialized_peak * 0.7, (lazy_peak, materialized_peak)


def test_reading_checks_the_metadata_size_before_openpyxl_parses_anything(
    tmp_path, monkeypatch
) -> None:
    import os
    import zipfile
    from io import BytesIO

    from app import spreadsheet_reader

    monkeypatch.setenv("INGESTION_MAX_OPEN_BYTES", "100000")
    base = _save(tmp_path, _simple(2)).read_bytes()
    padded = tmp_path / "pesada.xlsx"
    with (
        zipfile.ZipFile(BytesIO(base)) as original,
        zipfile.ZipFile(padded, "w", zipfile.ZIP_DEFLATED) as modified,
    ):
        for entry in original.infolist():
            modified.writestr(entry, original.read(entry))
        modified.writestr("xl/sharedStrings.xml", os.urandom(100_000).hex().encode())
    opened = []
    monkeypatch.setattr(
        spreadsheet_reader,
        "load_workbook",
        lambda *args, **kwargs: opened.append(args) or pytest.fail("parsed"),
    )

    with pytest.raises(IngestionLimitError) as excinfo:
        read_tables(padded, ".xlsx", "N-FP")

    assert excinfo.value.code == "workbook_metadata_limit"
    assert opened == []
