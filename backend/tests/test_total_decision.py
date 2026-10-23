from __future__ import annotations

import json
import sys
from pathlib import Path

import pytest
from openpyxl import load_workbook

from app.execution import validate_transition
from app.ingestion_budget import IngestionLimitError, budget_for
from app.pipeline import read_source, run_pipeline_batch
from app.upload_security import policy_for

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT))
from scripts.generate_total_decision_fixture import (  # noqa: E402
    DENIED_SENTINEL,
    SENSITIVE_SENTINEL,
    build_workbook,
)


class Repository:
    def __init__(self):
        self.state = "pending"

    def transition_execution(self, execution_id, target, reason):
        validate_transition(self.state, target)
        self.state = target

    def commit_pipeline(self, execution_id, result):
        self.transition_execution(execution_id, result["status"], result["status"])


def run(path):
    tables, issues = read_source(path, ".xlsx", "Total Decision")
    result = run_pipeline_batch(
        execution_id="synthetic-decision",
        inputs=[
            {
                "file_name": "synthetic.xlsx",
                "source": "Total Decision",
                "source_file_id": 1,
                "tables": tables,
                "read_issues": issues,
            }
        ],
        repository=Repository(),
        classified_at="2026-01-10T12:00:00+00:00",
    )
    return result


@pytest.fixture
def workbook(tmp_path):
    path = tmp_path / "total.xlsx"
    path.write_bytes(build_workbook())
    return path


def mutate(path, action):
    book = load_workbook(path)
    action(book.active)
    book.save(path)
    book.close()


def test_canonical_pipeline_preserves_identity_and_physical_provenance(workbook):
    result = run(workbook)
    assert result["status"] == "completed"
    assert len(result["normalized_records"]) == 4
    assert {r["row"] for r in result["normalized_records"]} == {4, 5, 6, 7}
    record = result["normalized_records"][0]
    assert record["source"] == "GMES/OQC"
    assert record["values"]["oqc_lot_id"] == "000000"
    assert record["values"]["oqc_id"] == "0000"
    assert record["values"]["workorder_id"] == "009990"
    assert record["values"]["source_organization_code"] == "SOURCE-SYN"
    assert "organization_code" not in record["values"]
    assert record["values"]["inspection_date"] == "2026-01-03"
    assert record["original_values"]["inspection_date"] == "03/01/2026"
    transform = next(
        t for t in record["transformations"] if t["field"] == "total_decision_result"
    )
    assert transform["adapter"] == "total_decision@1.0.0"
    assert transform["source_column"] == "AK:Total Decision / Result"
    text = json.dumps(result, default=str)
    for excluded in (SENSITIVE_SENTINEL, DENIED_SENTINEL, "NEVER-IMPORT-SUMMARY"):
        assert excluded not in text
    assert record["values"]["total_decision_remark_present"] is True


def test_edge_rows_rejected_independently_and_unknowns_preserved(tmp_path):
    path = tmp_path / "edge.xlsx"
    path.write_bytes(build_workbook(edge=True))
    result = run(path)
    assert result["status"] == "completed_with_errors"
    assert len(result["normalized_records"]) == 1
    assert {i["code"] for i in result["issues"]} >= {
        "invalid_quantity",
        "missing_identity",
        "formula_not_allowed",
        "unknown_state",
        "unparsed_date",
    }
    record = result["normalized_records"][0]
    assert record["values"]["total_decision_result"] == "SYN-UNKNOWN"
    assert record["values"]["production_date"] == "yyyy-MM-dd HH:mm"


@pytest.mark.parametrize(
    "cell,value",
    [
        ("AK1", "Wrong Group"),
        ("AL2", "Missing Date"),
        ("J1", None),
        ("AM2", "Unrelated User"),
    ],
)
def test_drift_blocks_before_records_are_imported(workbook, cell, value):
    mutate(workbook, lambda sheet: setattr(sheet[cell], "value", value))
    result = run(workbook)
    assert result["status"] == "validation_failed"
    assert not result["normalized_records"]
    assert not result["imported_records"]
    assert "contract_drift" in {i["code"] for i in result["issues"]}


def test_extra_column_and_sheet_are_explicitly_ignored(workbook):
    book = load_workbook(workbook)
    book.active["BW1"] = "Additional Field"
    book.active["BW4"] = DENIED_SENTINEL
    book.create_sheet("SYNTHETIC-PRIVATE-NAME")["A1"] = SENSITIVE_SENTINEL
    book.save(workbook)
    result = run(workbook)
    text = json.dumps(result, default=str)
    assert "SYNTHETIC-PRIVATE-NAME" not in text
    assert DENIED_SENTINEL not in text
    assert {i["code"] for i in result["issues"]} >= {
        "column_ignored",
        "sheet_ignored",
    }


def test_numeric_identifier_warns_and_absent_quantities_stay_null(workbook):
    def edit(sheet):
        sheet["BE4"] = 123
        sheet["AX4"] = None

    mutate(workbook, edit)
    result = run(workbook)
    first = result["normalized_records"][0]
    assert first["values"]["oqc_lot_id"] == "123"
    assert first["original_values"]["oqc_lot_id"] == 123
    assert first["values"]["hold_quantity"] is None
    assert "identifier_stored_as_number" in {i["code"] for i in result["issues"]}


def test_workorder_id_alone_does_not_become_workorder_number(workbook):
    mutate(workbook, lambda sheet: setattr(sheet["J4"], "value", None))
    result = run(workbook)
    record = result["normalized_records"][0]
    assert record["values"]["workorder_number"] is None
    assert record["values"]["workorder_id"] == "009990"
    assert "workorder_relationship_pending" in {i["code"] for i in result["issues"]}


def test_sensitive_formulas_never_escape_into_artifacts(workbook):
    mutate(
        workbook, lambda sheet: setattr(sheet["AN4"], "value", "=" + SENSITIVE_SENTINEL)
    )
    result = run(workbook)
    assert SENSITIVE_SENTINEL not in json.dumps(result, default=str)
    assert len(result["normalized_records"]) == 3


def test_policy_remains_xlsx_only_and_sheets_cannot_be_widened(monkeypatch):
    monkeypatch.setenv("UPLOAD_ALLOWED_EXTENSIONS", "csv,json,xlsx")
    monkeypatch.setenv("UPLOAD_ALLOWED_EXTENSIONS_TOTAL_DECISION", "csv,json,xlsx")
    monkeypatch.setenv("INGESTION_ALLOWED_SHEETS_TOTAL_DECISION", "*|Other")
    assert policy_for("Total Decision").allowed_extensions == {".xlsx"}
    assert budget_for("Total Decision").allowed_sheets == {"sheet1"}
    assert budget_for("Total Decision").selected_columns["sheet1"]


def test_row_budget_rejection_and_determinism(workbook, monkeypatch):
    first = run(workbook)
    assert first == run(workbook)
    monkeypatch.setenv("INGESTION_MAX_ROWS_TOTAL_DECISION", "4")
    with pytest.raises(IngestionLimitError, match="workbook_row_limit"):
        run(workbook)


def test_restricted_only_row_is_rejected_and_blank_row_is_ignored(workbook):
    def edit(sheet):
        sheet["AN9"] = SENSITIVE_SENTINEL

    mutate(workbook, edit)
    result = run(workbook)
    assert len(result["normalized_records"]) == 4
    assert len(result["imported_records"]) == 5
    assert SENSITIVE_SENTINEL not in json.dumps(result, default=str)
