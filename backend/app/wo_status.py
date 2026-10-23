from __future__ import annotations

import re
from collections import Counter
from collections.abc import Iterator
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Any

from openpyxl.utils.datetime import from_excel

from app.ingestion_budget import IngestionLimitError
from app.spreadsheet_reader import (
    IngestionRun,
    LazySheetRows,
    LocatedRow,
    SourceColumn,
    open_workbook,
)
from app.validation import SCHEMAS, _issue
from app.wo_status_spec import COMPOSITE_SOURCE, ColumnSpec, SheetSpec, load_spec

WO_STATUS_SCHEMA = SCHEMAS[COMPOSITE_SOURCE]
EXCEL_SERIAL_RANGE = (1, 2_958_465)
TEXT_DATE_FORMATS = (
    "%Y-%m-%d",
    "%d/%m/%Y",
    "%d-%m-%Y",
    "%Y-%m-%d %H:%M:%S",
    "%Y-%m-%dT%H:%M:%S",
    "%d/%m/%Y %H:%M:%S",
)
MONTHS = {
    name: index
    for index, name in enumerate(
        "jan feb mar apr may jun jul aug sep oct nov dec".split(), start=1
    )
}
# Dates written as ``05/Oct/2026 10:30:00`` by the source system (English month
# abbreviations, independent of the process locale).
NAMED_MONTH_DATE = re.compile(
    r"(?P<day>\d{1,2})/(?P<month>[A-Za-z]{3})/(?P<year>\d{4})"
    r"(?: (?P<h>\d{2}):(?P<m>\d{2}):(?P<s>\d{2}))?"
)
BLOCKED_REASON = "Aba bloqueada pelo contrato: a chave literal não foi confirmada"


class AdapterRow(LocatedRow):
    """Converted row that carries the issues found while reading its cells."""

    def __init__(self, values: list[Any], row_number: int) -> None:
        super().__init__(values, row_number)
        self.adapter_issues: list[dict[str, Any]] = []


class DriftedRows(list):
    """Placeholder rows of a sheet whose header no longer matches the contract."""

    def __init__(self, spec: SheetSpec) -> None:
        super().__init__()
        self.spec = spec
        self.stats: Counter[str] = Counter()


def _error(code: str, reason: str, column: str | None = None) -> dict[str, Any]:
    return _issue(code, reason, column=column)


def _warning(code: str, reason: str, column: str | None = None) -> dict[str, Any]:
    return _issue(code, reason, severity="warning", column=column)


def _is_formula(value: Any) -> bool:
    return isinstance(value, str) and value.lstrip().startswith("=")


def _identifier(value: Any, letter: str) -> tuple[Any, list[dict[str, Any]]]:
    if isinstance(value, bool):
        return str(value), []
    if isinstance(value, int):
        return str(value), [
            _warning(
                "identifier_stored_as_number",
                "Identificador armazenado como número; zeros à esquerda não "
                "podem ser recuperados",
                letter,
            )
        ]
    if isinstance(value, float):
        text = str(int(value)) if value.is_integer() else repr(value)
        return text, [
            _warning(
                "identifier_stored_as_number",
                "Identificador armazenado como número; zeros à esquerda não "
                "podem ser recuperados",
                letter,
            )
        ]
    text = str(value).strip()
    return (text or None), []


def _quantity(value: Any, letter: str, integer: bool) -> tuple[Any, list[dict]]:
    if isinstance(value, bool):
        return None, [_error("invalid_quantity", "Quantidade inválida", letter)]
    try:
        number = Decimal(str(value).strip())
    except InvalidOperation:
        return None, [_error("invalid_quantity", "Quantidade inválida", letter)]
    if (
        not number.is_finite()
        or number < 0
        or (integer and number != number.to_integral())
    ):
        return None, [_error("invalid_quantity", "Quantidade inválida", letter)]
    return (int(number) if number == number.to_integral() else float(number)), []


def _temporal(value: Any, letter: str, kind: str) -> tuple[Any, list[dict]]:
    if isinstance(value, datetime | date):
        return value, []
    if isinstance(value, int | float) and not isinstance(value, bool):
        low, high = EXCEL_SERIAL_RANGE
        if low <= value < high:
            return from_excel(value), []
        return None, [_error("invalid_date", "Data inválida", letter)]
    text = str(value).strip()
    named = NAMED_MONTH_DATE.fullmatch(text)
    if named and named["month"].casefold() in MONTHS:
        day, month, year = (
            int(named["day"]),
            MONTHS[named["month"].casefold()],
            int(named["year"]),
        )
        hour, minute, second = (int(named[k] or 0) for k in ("h", "m", "s"))
        try:
            return datetime(year, month, day, hour, minute, second), []
        except ValueError:
            return None, [_error("invalid_date", "Data inválida", letter)]
    for date_format in TEXT_DATE_FORMATS:
        try:
            return datetime.strptime(text, date_format), []
        except ValueError:
            continue
    return text, [
        _warning(
            "unparsed_date",
            "Data em formato não reconhecido; valor original preservado",
            letter,
        )
    ]


def convert(column: ColumnSpec, value: Any) -> tuple[Any, list[dict[str, Any]]]:
    """Convert one literal cell; formulas never become source data."""
    if value is None or (isinstance(value, str) and not value.strip()):
        return None, []
    if _is_formula(value):
        return None, [
            _error(
                "formula_not_allowed",
                "Célula calculada não pode ser fonte primária",
                column.letter,
            )
        ]
    kind = column.kind
    if kind == "id":
        return _identifier(value, column.letter)
    if kind in {"qty", "qty_int"}:
        return _quantity(value, column.letter, kind == "qty_int")
    if kind in {"date", "datetime"}:
        return _temporal(value, column.letter, kind)
    if isinstance(value, str):
        return value.strip(), []
    return value, []


class AdapterRows(LazySheetRows):
    """Canonical rows of one authorized sheet, streamed with its contract rules."""

    def __init__(
        self, path: Path, spec: SheetSpec, adapter: str, run: IngestionRun
    ) -> None:
        super().__init__(path, spec.name, spec.header_row - 1, run)
        self.spec = spec
        self.adapter = adapter
        self.stats: Counter[str] = Counter()
        self.headers = [
            SourceColumn(
                f"{column.letter}:{column.header}", column.field, column.letter, adapter
            )
            for column in spec.columns
        ]

    def _identity_missing(self, present: dict[str, bool]) -> bool:
        spec = self.spec
        if spec.identity_all and not all(present.get(f) for f in spec.identity_all):
            return True
        return bool(spec.identity_any) and not any(
            present.get(f) for f in spec.identity_any
        )

    def _placeholder(self, number: int, code: str, reason: str) -> AdapterRow:
        row = AdapterRow([None] * len(self.spec.columns), number)
        row.adapter_issues.append(_error(code, reason))
        return row

    def _convert(
        self, number: int, values: tuple
    ) -> tuple[AdapterRow, dict[str, bool]]:
        converted: list[Any] = []
        issues: list[dict[str, Any]] = []
        present: dict[str, bool] = {}
        for column in self.spec.columns:
            raw = values[column.index - 1] if column.index <= len(values) else None
            value, cell_issues = convert(column, raw)
            converted.append(value)
            issues.extend(cell_issues)
            present[column.field] = value is not None
        row = AdapterRow(converted, number)
        row.adapter_issues = issues
        return row, present

    def __iter__(self) -> Iterator[LocatedRow]:
        spec = self.spec
        budget = self.run.budget
        stats: Counter[str] = Counter()
        self.stats = stats
        gap: list[int] = []
        cells = 0
        width = len(spec.columns)
        try:
            with open_workbook(self.path, budget) as workbook:
                worksheet = workbook[spec.name]
                for number, values in enumerate(
                    worksheet.iter_rows(
                        min_row=spec.data_start_row,
                        max_col=spec.max_column,
                        values_only=True,
                    ),
                    spec.data_start_row,
                ):
                    if number > budget.max_rows:
                        raise IngestionLimitError(
                            "workbook_row_limit",
                            limit=budget.max_rows,
                            observed=number,
                            sheet=spec.name,
                        )
                    cells += width
                    if cells > budget.max_cells:
                        raise IngestionLimitError(
                            "workbook_cell_limit",
                            limit=budget.max_cells,
                            observed=cells,
                            sheet=spec.name,
                        )
                    estimate = budget.minimum_memory(number, cells)
                    if estimate > budget.max_memory_bytes:
                        raise IngestionLimitError(
                            "workbook_memory_limit",
                            limit=budget.max_memory_bytes,
                            observed=estimate,
                            sheet=spec.name,
                        )
                    self.run.tick()
                    stats["scanned"] += 1
                    row, present = self._convert(number, values)
                    blank = not any(present.values()) and not row.adapter_issues
                    if spec.window_mode == "substantive":
                        if not present.get(spec.window_field or ""):
                            stats["ignored_non_substantive"] += 1
                            continue
                        if spec.blocked:
                            stats["rejected_blocked"] += 1
                            yield self._placeholder(
                                number, "contract_blocked", BLOCKED_REASON
                            )
                            continue
                    elif blank:
                        if spec.window_mode == "last_literal":
                            gap.append(number)
                        stats["ignored_blank"] += 1
                        continue
                    if spec.window_mode == "last_literal" and present.get(
                        spec.window_field or ""
                    ):
                        for gap_number in gap:
                            stats["ignored_blank"] -= 1
                            stats["rejected_gap"] += 1
                            yield self._placeholder(
                                gap_number,
                                "gap_in_active_region",
                                "Linha vazia dentro da região ativa da aba",
                            )
                        gap.clear()
                    if self._identity_missing(present):
                        row.adapter_issues.append(
                            _error(
                                "missing_identity",
                                "Informe os identificadores exigidos pelo contrato",
                            )
                        )
                    if any(i["severity"] == "error" for i in row.adapter_issues):
                        stats["rejected_adapter"] += 1
                    stats["emitted"] += 1
                    yield row
                stats["ignored_tail"] = len(gap)
                stats["ignored_blank"] -= len(gap)
        except IngestionLimitError:
            raise
        except Exception as exc:
            raise IngestionLimitError("unreadable_workbook", sheet=spec.name) from exc


class WoStatusImporter:
    """Reads the authorized sheets of a WO Status workbook as canonical tables."""

    def read(
        self, path: Path, extension: str, source: str
    ) -> tuple[list[tuple[str, list[str], Any]], list[dict[str, Any]]]:
        if extension != ".xlsx":
            raise ValueError("A fonte WO Status exige arquivo XLSX")
        spec = load_spec()
        run = IngestionRun(source)
        tables: list[tuple[str, list[str], Any]] = []
        issues: list[dict[str, Any]] = []
        seen: set[str] = set()
        with open_workbook(path, run.budget) as workbook:
            for name in workbook.sheetnames:
                sheet = spec.sheet(name)
                if sheet is None:
                    issues.append(
                        _issue(
                            "sheet_ignored",
                            "Aba fora do contrato; não foi ingerida",
                            severity="warning",
                            sheet=name,
                        )
                    )
                    continue
                seen.add(sheet.name)
                if not run.budget.authorizes(name):
                    issues.append(
                        _issue(
                            "sheet_not_authorized",
                            "Aba do contrato não autorizada neste ambiente",
                            severity="warning",
                            sheet=name,
                        )
                    )
                    continue
                drift = self._header_drift(workbook[name], sheet)
                if drift:
                    issues.extend({**item, "sheet": name} for item in drift)
                    tables.append((name, [], DriftedRows(sheet)))
                    continue
                rows = AdapterRows(path, sheet, spec.label, run)
                tables.append((name, rows.headers, rows))
        for sheet in spec.sheets:
            if sheet.name not in seen and run.budget.authorizes(sheet.name):
                issues.append(
                    _issue(
                        "sheet_missing",
                        "Aba do contrato ausente no arquivo",
                        severity="warning",
                        sheet=sheet.name,
                    )
                )
        return tables, issues

    @staticmethod
    def _header_drift(worksheet: Any, sheet: SheetSpec) -> list[dict[str, Any]]:
        header = next(
            worksheet.iter_rows(
                min_row=sheet.header_row,
                max_row=sheet.header_row,
                max_col=sheet.max_column,
                values_only=True,
            ),
            (),
        )
        drift = []
        for column in sheet.columns:
            cell = header[column.index - 1] if column.index <= len(header) else None
            found = re.sub(r"\s+", " ", str(cell or "")).strip().casefold()
            if found != column.header.casefold():
                item = _issue(
                    "contract_drift",
                    f"Cabeçalho da coluna {column.letter} diverge do contrato",
                    column=column.letter,
                    row=sheet.header_row,
                )
                drift.append(item)
        return drift


def is_composite(source: str) -> bool:
    return source == COMPOSITE_SOURCE


def expand_composite(item: dict[str, Any]) -> list[dict[str, Any]]:
    """Split one WO Status file into one pipeline input per authorized sheet.

    Each sheet keeps the file identifier but carries its logical source and the
    adapter schema, so records enter the existing N-FP, OWM and TMS flows.
    """
    sheets = [(sheet, headers, rows) for sheet, headers, rows in item["tables"]]
    if not sheets:
        return [{**item, "schema": WO_STATUS_SCHEMA}]
    by_sheet: dict[str, list[dict[str, Any]]] = {}
    general: list[dict[str, Any]] = []
    names = {sheet for sheet, _, _ in sheets}
    for issue in item["read_issues"]:
        if issue.get("sheet") in names:
            by_sheet.setdefault(issue["sheet"], []).append(issue)
        else:
            general.append(issue)
    expanded = []
    for position, table in enumerate(sheets):
        sheet, _, rows = table
        expanded.append(
            {
                **item,
                "source": rows.spec.logical_source,
                "tables": [table],
                "read_issues": [
                    *(general if position == 0 else []),
                    *by_sheet.get(sheet, []),
                ],
                "schema": WO_STATUS_SCHEMA,
            }
        )
    return expanded
