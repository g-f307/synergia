from __future__ import annotations

import logging
import zipfile
from collections.abc import Iterator
from contextlib import contextmanager
from itertools import islice
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from app.ingestion_budget import IngestionBudget, IngestionLimitError, budget_for
from app.observability.telemetry import safe_log
from app.xlsx_package import check_open_parts

HEADER_SCAN_ROWS = 200
CLOCK_CHECK_INTERVAL = 1000


@contextmanager
def open_workbook(
    path: Path | str, budget: IngestionBudget | None = None
) -> Iterator[Any]:
    """Open a workbook read-only from a stream, whatever the file suffix is.

    Stored uploads keep a neutral suffix, which openpyxl refuses when given a
    path, so the file is always handed over as an open binary stream.
    """
    with Path(path).open("rb") as stream:
        try:
            if budget is not None:
                # openpyxl loads shared strings and styles whole on opening, so
                # their size is checked before it parses anything.
                with zipfile.ZipFile(stream) as archive:
                    check_open_parts(archive, budget)
                stream.seek(0)
            workbook = load_workbook(stream, read_only=True, data_only=False)
        except IngestionLimitError:
            raise
        except Exception as exc:
            raise IngestionLimitError("unreadable_workbook") from exc
        try:
            yield workbook
        finally:
            workbook.close()


class SourceColumn(str):
    """Header of an adapter column.

    Its text is the origin (``"Y:Demand ID"``) so provenance keeps the original
    column; ``canonical`` is the field the pipeline works with.
    """

    canonical: str
    letter: str
    adapter: str | None

    def __new__(
        cls, label: str, canonical: str, letter: str, adapter: str | None = None
    ) -> SourceColumn:
        column = super().__new__(cls, label)
        column.canonical = canonical
        column.letter = letter
        column.adapter = adapter
        return column


class LocatedRow(list[Any]):
    """Spreadsheet row that retains its physical position after preambles."""

    def __init__(self, values: list[Any], row_number: int) -> None:
        super().__init__(values)
        self.row_number = row_number


class IngestionRun:
    """Shared budget, clock and progress counters for one file's processing."""

    def __init__(
        self,
        source: str,
        budget: IngestionBudget | None = None,
        *,
        execution_id: str | None = None,
    ) -> None:
        self.source = source
        self.budget = budget or budget_for(source)
        self.execution_id = execution_id
        self.clock = self.budget.start()
        self.phase = "reading"
        self.rows_by_phase: dict[str, int] = {}

    def restart(self) -> None:
        """Start the processing window once the file is about to be consumed."""
        self.clock = self.budget.start()
        self.rows_by_phase = {}

    def begin(self, phase: str) -> None:
        self.phase = phase
        self.rows_by_phase.setdefault(phase, 0)
        self.clock.check()

    def tick(self) -> None:
        before = self.rows_by_phase.get(self.phase, 0)
        after = before + 1
        self.rows_by_phase[self.phase] = after
        if after // CLOCK_CHECK_INTERVAL != before // CLOCK_CHECK_INTERVAL:
            self.clock.check()
        interval = self.budget.progress_interval
        if after // interval != before // interval:
            safe_log(
                logging.INFO,
                "ingestion.progress",
                source=self.source,
                operation=self.phase,
                processed_count=after,
                duration_ms=int(self.clock.elapsed * 1000),
            )

    @property
    def rows_read(self) -> int:
        return max(self.rows_by_phase.values(), default=0)


class LazySheetRows:
    """Re-iterable rows of one worksheet, streamed from the stored file.

    Each pass reopens the workbook in read-only mode, so no pass keeps the sheet
    in memory. Budget limits are enforced again while reading as a defence in
    depth against files that changed after admission.
    """

    def __init__(
        self, path: Path, sheet: str, header_index: int, run: IngestionRun
    ) -> None:
        self.path = Path(path)
        self.sheet = sheet
        self.header_index = header_index
        self.header_row = header_index + 1
        self.run = run
        self._length: int | None = None

    def __iter__(self) -> Iterator[LocatedRow]:
        budget = self.run.budget
        cells = 0
        try:
            with open_workbook(self.path, budget) as workbook:
                worksheet = workbook[self.sheet]
                for number, values in enumerate(
                    worksheet.iter_rows(values_only=True), 1
                ):
                    if number <= self.header_row:
                        continue
                    if number > budget.max_rows:
                        raise IngestionLimitError(
                            "workbook_row_limit",
                            limit=budget.max_rows,
                            observed=number,
                            sheet=self.sheet,
                        )
                    width = len(values)
                    if width > budget.max_columns:
                        raise IngestionLimitError(
                            "workbook_column_limit",
                            limit=budget.max_columns,
                            observed=width,
                            sheet=self.sheet,
                        )
                    cells += width
                    if cells > budget.max_cells:
                        raise IngestionLimitError(
                            "workbook_cell_limit",
                            limit=budget.max_cells,
                            observed=cells,
                            sheet=self.sheet,
                        )
                    self._check_memory(budget.minimum_memory(number, cells))
                    self.run.tick()
                    yield LocatedRow(list(values), number)
        except IngestionLimitError:
            raise
        except Exception as exc:
            raise IngestionLimitError("unreadable_workbook", sheet=self.sheet) from exc

    def _check_memory(self, estimate: int) -> None:
        limit = self.run.budget.max_memory_bytes
        if estimate > limit:
            raise IngestionLimitError(
                "workbook_memory_limit",
                limit=limit,
                observed=estimate,
                sheet=self.sheet,
            )

    def count(self) -> int:
        """Count rows with one extra pass.

        ``__len__`` is deliberately not defined: ``list(rows)`` would call it
        and silently read the whole sheet twice.
        """
        if self._length is None:
            self._length = sum(1 for _ in self)
        return self._length

    def __bool__(self) -> bool:
        return next(iter(self), None) is not None

    def __getitem__(self, index: int | slice):
        if isinstance(index, slice):
            return list(islice(self, index.start, index.stop, index.step))
        if index < 0:
            raise IndexError("índice negativo não suportado em leitura incremental")
        for row in islice(self, index, index + 1):
            return row
        raise IndexError(index)
