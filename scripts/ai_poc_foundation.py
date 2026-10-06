"""Shared, local-only foundation for the Demo Day AI PoCs.

This module is deliberately outside ``backend/app`` and ``web``.  It handles
synthetic files, deterministic baselines and an optional loopback Ollama
runtime without changing the production SYNERGIA flow.
"""

from __future__ import annotations

import csv
import hashlib
import ipaddress
import json
import math
import os
import statistics
import time
import tracemalloc
import urllib.error
import urllib.request
import uuid
from collections.abc import Mapping
from dataclasses import dataclass
from pathlib import Path
from typing import Any, Protocol
from urllib.parse import urlsplit

from jsonschema import Draft202012Validator, FormatChecker
from openpyxl import load_workbook


class PoCFoundationError(Exception):
    """Expected, actionable error in local PoC execution."""


class LocalRuntimeTimeout(PoCFoundationError):
    """The configured local generation deadline elapsed."""


SENSITIVE_MARKERS = (
    "password",
    "authorization",
    "bearer ",
    "secret",
    "token",
    "cookie",
    "api_key",
)


class ModelRuntime(Protocol):
    def generate(self, prompt: str, *, model: str) -> tuple[str, int]: ...


@dataclass(frozen=True)
class DatasetSpec:
    required_columns: tuple[str, ...]
    column_types: Mapping[str, str]
    domains: Mapping[str, frozenset[str]]
    max_bytes: int = 25 * 1024 * 1024


@dataclass(frozen=True)
class LocalModelConfig:
    model: str
    endpoint: str = "http://127.0.0.1:11434/api/generate"
    timeout_seconds: float = 120.0

    def __post_init__(self) -> None:
        _validate_loopback_url(self.endpoint)

    @classmethod
    def from_environment(cls) -> LocalModelConfig:
        endpoint = os.getenv("SYNERGIA_POC_MODEL_ENDPOINT", cls.endpoint)
        model = os.getenv("SYNERGIA_POC_MODEL", "qwen2.5:3b-instruct-q4_K_M")
        if not model or any(char in model for char in "\r\n"):
            raise PoCFoundationError("SYNERGIA_POC_MODEL inválido")
        return cls(model=model, endpoint=endpoint)


class OllamaRuntime:
    def __init__(
        self,
        config: LocalModelConfig,
        *,
        options: dict[str, Any] | None = None,
        output_schema: dict[str, Any] | None = None,
    ) -> None:
        _validate_loopback_url(config.endpoint)
        self.config = config
        self.options = dict(options or {})
        self.output_schema = output_schema
        self.response_model: str | None = None
        self.response_metrics: dict[str, int] = {}

    def generate(self, prompt: str, *, model: str) -> tuple[str, int]:
        document = {"model": model, "prompt": prompt, "stream": False}
        if self.options:
            document["options"] = self.options
        if self.output_schema is not None:
            document["format"] = self.output_schema
        body = json.dumps(document).encode()
        request = urllib.request.Request(
            self.config.endpoint,
            data=body,
            headers={"Content-Type": "application/json"},
            method="POST",
        )
        try:
            opener = urllib.request.build_opener(_NoRedirectHandler())
            with opener.open(request, timeout=self.config.timeout_seconds) as response:
                payload = json.loads(response.read().decode("utf-8"))
        except PoCFoundationError:
            raise
        except TimeoutError as exc:
            raise LocalRuntimeTimeout("runtime local excedeu o tempo limite") from exc
        except urllib.error.URLError as exc:
            if isinstance(exc.reason, TimeoutError):
                raise LocalRuntimeTimeout(
                    "runtime local excedeu o tempo limite"
                ) from exc
            raise PoCFoundationError("runtime local indisponível") from exc
        except json.JSONDecodeError as exc:
            raise PoCFoundationError("runtime local indisponível") from exc
        output = payload.get("response")
        if not isinstance(output, str):
            raise PoCFoundationError("runtime local retornou resposta sem JSON textual")
        self.response_model = payload.get("model")
        self.response_metrics = {
            key: payload[key]
            for key in (
                "load_duration",
                "total_duration",
                "prompt_eval_count",
                "prompt_eval_duration",
                "eval_count",
                "eval_duration",
            )
            if isinstance(payload.get(key), int) and payload[key] >= 0
        }
        return output, int(payload.get("eval_count") or 0)


def _validate_loopback_url(endpoint: str) -> None:
    try:
        parsed = urlsplit(endpoint)
        hostname = parsed.hostname
        port = parsed.port
    except ValueError as exc:
        raise PoCFoundationError("endpoint local inválido") from exc
    if (
        parsed.scheme != "http"
        or not hostname
        or parsed.username is not None
        or parsed.password is not None
        or parsed.query
        or parsed.fragment
        or port is None
    ):
        raise PoCFoundationError("endpoint deve ser HTTP loopback sem credenciais")
    if hostname.lower() == "localhost":
        return
    try:
        address = ipaddress.ip_address(hostname)
    except ValueError as exc:
        raise PoCFoundationError("endpoint deve apontar para loopback") from exc
    if not address.is_loopback:
        raise PoCFoundationError("endpoint deve apontar para loopback")


class _NoRedirectHandler(urllib.request.HTTPRedirectHandler):
    def redirect_request(self, request, fp, code, msg, headers, new_url):
        raise PoCFoundationError("redirecionamento do runtime local bloqueado")


def _coerce(value: Any, expected: str, *, column: str, row: int) -> Any:
    if expected == "string":
        if not isinstance(value, str) or not value.strip():
            raise PoCFoundationError(f"linha {row}: coluna {column} deve ser texto")
        return value.strip()
    if expected == "integer":
        try:
            converted = int(value)
        except (TypeError, ValueError) as exc:
            raise PoCFoundationError(
                f"linha {row}: coluna {column} deve ser inteira"
            ) from exc
        return converted
    if expected == "number":
        try:
            return float(value)
        except (TypeError, ValueError) as exc:
            raise PoCFoundationError(
                f"linha {row}: coluna {column} deve ser numérica"
            ) from exc
    raise PoCFoundationError(f"tipo não suportado para coluna {column}: {expected}")


def validate_records(
    records: list[dict[str, Any]], spec: DatasetSpec
) -> list[dict[str, Any]]:
    if not records:
        raise PoCFoundationError("arquivo sem registros")
    actual = set(records[0])
    missing = set(spec.required_columns) - actual
    if missing:
        raise PoCFoundationError(
            f"colunas obrigatórias ausentes: {', '.join(sorted(missing))}"
        )
    unexpected = actual - set(spec.column_types)
    if unexpected:
        raise PoCFoundationError(
            f"colunas não previstas: {', '.join(sorted(unexpected))}"
        )
    normalized: list[dict[str, Any]] = []
    for index, record in enumerate(records, start=2):
        item = {
            key: _coerce(record.get(key), kind, column=key, row=index)
            for key, kind in spec.column_types.items()
        }
        for column, allowed in spec.domains.items():
            if item[column] not in allowed:
                raise PoCFoundationError(f"linha {index}: domínio inválido em {column}")
        normalized.append(item)
    return normalized


def _read_csv(path: Path) -> list[dict[str, Any]]:
    with path.open(newline="", encoding="utf-8-sig") as handle:
        return list(csv.DictReader(handle))


def _read_xlsx(path: Path) -> list[dict[str, Any]]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        sheet = workbook.active
        rows = list(sheet.iter_rows(values_only=True))
    finally:
        workbook.close()
    if not rows:
        return []
    headers = [str(value).strip() if value is not None else "" for value in rows[0]]
    return [
        dict(zip(headers, row, strict=False))
        for row in rows[1:]
        if any(value is not None for value in row)
    ]


def ingest(path: Path, spec: DatasetSpec) -> list[dict[str, Any]]:
    path = path.resolve()
    if path.suffix.lower() not in {".csv", ".xlsx"}:
        raise PoCFoundationError("extensão inválida; use CSV ou XLSX")
    if not path.is_file():
        raise PoCFoundationError("arquivo de entrada não encontrado")
    if path.stat().st_size > spec.max_bytes:
        raise PoCFoundationError("arquivo excede o tamanho máximo permitido")
    records = _read_csv(path) if path.suffix.lower() == ".csv" else _read_xlsx(path)
    return validate_records(records, spec)


def load_prompt(path: Path) -> str:
    prompt = path.read_text(encoding="utf-8")
    if not prompt.strip() or len(prompt) > 32_000:
        raise PoCFoundationError("prompt vazio ou acima do limite")
    return prompt


def validate_output(payload: Any, schema_path: Path) -> None:
    schema = json.loads(schema_path.read_text(encoding="utf-8"))
    errors = sorted(
        Draft202012Validator(schema, format_checker=FormatChecker()).iter_errors(
            payload
        ),
        key=lambda error: list(error.path),
    )
    if errors:
        location = ".".join(str(part) for part in errors[0].path) or "$"
        raise PoCFoundationError(
            f"saída fora do esquema em {location}: {errors[0].message}"
        )


def _reject_sensitive_content(value: Any, path: str = "$") -> None:
    if isinstance(value, Mapping):
        for key, item in value.items():
            key_text = str(key).lower()
            if any(marker in key_text for marker in SENSITIVE_MARKERS):
                raise PoCFoundationError(f"conteúdo sensível na saída: {path}.{key}")
            _reject_sensitive_content(item, f"{path}.{key}")
    elif isinstance(value, list):
        for index, item in enumerate(value):
            _reject_sensitive_content(item, f"{path}[{index}]")
    elif isinstance(value, str):
        lowered = value.lower()
        if any(marker in lowered for marker in SENSITIVE_MARKERS):
            raise PoCFoundationError(f"conteúdo sensível na saída: {path}")


def run_baseline(kind: str, records: list[dict[str, Any]]) -> dict[str, Any]:
    """Run the deterministic comparator on the same normalized records."""
    if kind == "quality":
        diagnoses = [
            {
                "code": f"STATUS_{item['status'].upper()}",
                "severity": "error" if item["status"] == "error" else "warning",
                "message": f"Registro {item['record_id']} requer análise.",
            }
            for item in records
            if item["status"] != "valid"
        ]
        return {"schema_version": "1.0.0", "diagnoses": diagnoses}
    if kind == "operational":
        findings = [
            {
                "finding_id": item["record_id"],
                "severity": "info" if item["status"] == "valid" else "medium",
                "message": f"Registro {item['record_id']} em estado {item['status']}.",
            }
            for item in records
        ]
        return {
            "schema_version": "1.0.0",
            "summary": f"{len(records)} registros processados.",
            "findings": findings,
        }
    raise PoCFoundationError(f"PoC desconhecida: {kind}")


def normalized_input_hash(records: list[dict[str, Any]]) -> str:
    context = json.dumps(records, ensure_ascii=False, sort_keys=True)
    return hashlib.sha256(context.encode("utf-8")).hexdigest()


@dataclass(frozen=True)
class ExecutionMetrics:
    execution_id: str
    model: str
    input_sha256: str
    elapsed_ms: float
    generated_tokens: int
    tokens_per_second: float
    peak_memory_bytes: int


def execute_case(
    records: list[dict[str, Any]],
    prompt_path: Path,
    schema_path: Path,
    runtime: ModelRuntime,
    config: LocalModelConfig,
    *,
    output_dir: Path,
) -> tuple[dict[str, Any], ExecutionMetrics]:
    prompt = load_prompt(prompt_path)
    context = json.dumps(records, ensure_ascii=False, sort_keys=True)
    input_sha256 = normalized_input_hash(records)
    tracemalloc.start()
    started = time.perf_counter()
    raw, generated_tokens = runtime.generate(
        f"{prompt}\nDADOS:\n{context}", model=config.model
    )
    elapsed_ms = (time.perf_counter() - started) * 1000
    _, peak_memory = tracemalloc.get_traced_memory()
    tracemalloc.stop()
    try:
        output = json.loads(raw)
    except json.JSONDecodeError as exc:
        raise PoCFoundationError("modelo não produziu JSON válido") from exc
    validate_output(output, schema_path)
    _reject_sensitive_content(output)
    execution_id = uuid.uuid4().hex
    output_dir.mkdir(parents=True, exist_ok=True)
    (output_dir / f"{execution_id}.json").write_text(
        json.dumps(
            {"execution_id": execution_id, "output": output},
            ensure_ascii=False,
            indent=2,
        )
        + "\n",
        encoding="utf-8",
    )
    metrics = ExecutionMetrics(
        execution_id,
        config.model,
        input_sha256,
        round(elapsed_ms, 3),
        generated_tokens,
        round(generated_tokens / (elapsed_ms / 1000), 3) if elapsed_ms else 0.0,
        peak_memory,
    )
    (output_dir / f"{execution_id}.metrics.json").write_text(
        json.dumps(metrics.__dict__, indent=2) + "\n", encoding="utf-8"
    )
    return output, metrics


def summarize_metrics(metrics: list[ExecutionMetrics]) -> dict[str, float]:
    if not metrics:
        raise PoCFoundationError("nenhuma métrica para resumir")
    latencies = sorted(item.elapsed_ms for item in metrics)
    p95_index = min(len(latencies) - 1, max(0, math.ceil(len(latencies) * 0.95) - 1))
    return {
        "mean_latency_ms": round(statistics.mean(latencies), 3),
        "p95_latency_ms": round(latencies[p95_index], 3),
        "mean_tokens_per_second": round(
            statistics.mean(item.tokens_per_second for item in metrics), 3
        ),
        "peak_memory_bytes": max(item.peak_memory_bytes for item in metrics),
    }
