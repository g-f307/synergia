"""Generate invented data using only the published Total Decision structure."""

from __future__ import annotations

import argparse
import sys
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.utils.cell import column_index_from_string

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "backend"))
sys.path.insert(0, str(ROOT))

from app.total_decision_spec import load_spec  # noqa: E402
from scripts.generate_homologation_fixture import _stable_xlsx  # noqa: E402

SENSITIVE_SENTINEL = "SYNTHETIC-RESTRICTED-CONTENT"
DENIED_SENTINEL = "SYNTHETIC-DENIED-CONTENT"


def build_workbook(rows: int = 4, *, edge: bool = False) -> bytes:
    if rows < 1:
        raise ValueError("rows deve ser positivo")
    spec = load_spec()
    book = Workbook()
    book.properties.created = datetime(2026, 1, 1)
    book.properties.modified = datetime(2026, 1, 1)
    sheet = book.active
    sheet.title = spec.sheet.name
    for column in spec.sheet.columns:
        if len(column.header_path) == 1:
            sheet.cell(1, column.index, column.header_path[0])
            sheet.merge_cells(
                start_row=1,
                end_row=2,
                start_column=column.index,
                end_column=column.index,
            )
        else:
            sheet.cell(2, column.index, column.header_path[1])
    for span, label in (
        ("AK1:AN1", "Total Decision"),
        ("AP1:AR1", "ERP Upload"),
        ("AS1:AW1", "Approval"),
    ):
        sheet.merge_cells(span)
        sheet[span.split(":")[0]] = label
    for letter, label in (
        ("Z", "MANUAL_OWNER"),
        ("BG", "Inspection ID"),
        ("BH", "WOID"),
        ("AJ", "C"),
        ("AO", "C"),
        ("L", "Not in 6 Months"),
        ("BN", "Not in 6 Months"),
    ):
        sheet[f"{letter}1"] = label
    sheet["A3"] = "SYNTHETIC TOTALS"
    sheet["J3"] = "NEVER-IMPORT-SUMMARY"
    sheet["BE3"] = "NEVER-IMPORT-SUMMARY"
    sheet["AX3"] = 999999
    for index in range(rows):
        number = index + 4
        values = {
            "A": "2026-01-02",
            "B": "03/01/2026",
            "C": f"000{index}",
            "D": "LINE-SYN",
            "F": "MODEL-SYN",
            "G": "SFX",
            "J": f"SYN-WO-{index:04d}",
            "AK": "approved",
            "AL": "2026-01-03 12:30:00",
            "AM": SENSITIVE_SENTINEL,
            "AN": SENSITIVE_SENTINEL,
            "AP": "completed",
            "AQ": "2026-01-04 12:30:00",
            "AR": SENSITIVE_SENTINEL,
            "AX": 2,
            "AY": 5,
            "AZ": 1,
            "BA": 3,
            "BE": f"00000{index}",
            "BF": f"00999{index}",
            "BM": "SOURCE-SYN",
            "Z": DENIED_SENTINEL,
            "BG": DENIED_SENTINEL,
        }
        if edge:
            if index == 0:
                values.update(A="yyyy-MM-dd HH:mm", AK="SYN-UNKNOWN")
            elif index == 1:
                values["BA"] = -1
            elif index == 2:
                values["BE"] = None
            elif index == 3:
                values["AX"] = "=1+1"
        for letter, value in values.items():
            sheet.cell(number, column_index_from_string(letter), value)
    return _stable_xlsx(book)


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("output", type=Path)
    parser.add_argument("--rows", type=int, default=4)
    parser.add_argument("--edge", action="store_true")
    args = parser.parse_args()
    args.output.write_bytes(build_workbook(args.rows, edge=args.edge))


if __name__ == "__main__":
    main()
