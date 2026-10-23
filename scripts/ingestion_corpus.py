"""Count-only adapter evidence and frozen, entirely invented corpus inputs."""

from __future__ import annotations

import hashlib
import json
import sys
from collections import Counter
from io import BytesIO
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
sys.path.insert(0, str(ROOT / "backend"))

from scripts.generate_homologation_fixture import _stable_xlsx  # noqa: E402
from scripts.generate_total_decision_fixture import (  # noqa: E402
    build_workbook as decision,
)
from scripts.generate_wo_status_fixture import build_workbook as workorder  # noqa: E402

from app.pipeline import read_source, run_pipeline_batch  # noqa: E402
from app.total_decision_spec import load_spec as decision_spec  # noqa: E402
from app.upload_security import inspect_content, policy_for  # noqa: E402
from app.wo_status_report import _ValidationOnlyRepository  # noqa: E402
from app.wo_status_spec import load_spec as workorder_spec  # noqa: E402

CORPUS = ROOT / "data/synthetic/ingestion-v1"
VERSION = "1.0.0"
CLASSIFIED_AT = "2026-01-10T00:00:00+00:00"


def canonical(value: object) -> bytes:
    return (
        json.dumps(
            value,
            sort_keys=True,
            indent=2,
            allow_nan=False,
            ensure_ascii=False,
            default=str,
        )
        + "\n"
    ).encode()


def digest(content: bytes) -> str:
    return hashlib.sha256(content).hexdigest()


def specs(source: str):
    if source == "WO Status":
        return workorder_spec().sheets
    if source == "Total Decision":
        return (decision_spec().sheet,)
    raise ValueError("Fonte fora do corpus")


def execute(
    path: Path,
    source: str,
    *,
    repository=None,
    execution_id="corpus-validation",
    source_file_id=0,
):
    """Use the real admission and pipeline; never return unknown sheet names."""
    policy = policy_for(source)
    if path.stat().st_size > policy.max_bytes:
        from app.upload_security import UnsafeUpload

        raise UnsafeUpload("file_too_large")
    inspect_content(path, ".xlsx", policy)
    tables, issues = read_source(path, ".xlsx", source)
    result = run_pipeline_batch(
        execution_id=execution_id,
        inputs=[
            {
                "source": source,
                "file_name": "synthetic.xlsx",
                "source_file_id": source_file_id,
                "tables": tables,
                "read_issues": issues,
            }
        ],
        repository=repository or _ValidationOnlyRepository(),
        classified_at=CLASSIFIED_AT,
    )
    return summarize(source, tables, result), result


def summarize(source, tables, result):
    by_name = {name: rows for name, _, rows in tables}
    imported = Counter(r["sheet"] for r in result["imported_records"])
    normalized = Counter(r["sheet"] for r in result["normalized_records"])
    sheets = []
    for spec in specs(source):
        rows = by_name.get(spec.name)
        stats = getattr(rows, "stats", Counter())
        issues = [i for i in result["issues"] if i.get("sheet") == spec.name]
        rejected = {
            i["row"]
            for i in issues
            if i["severity"] == "error" and (i.get("row") or 0) >= spec.data_start_row
        }
        duplicates = {
            i["row"]
            for i in issues
            if i["code"] in {"duplicate_row", "duplicate_serial"}
        }
        ignored = sum(
            stats[k]
            for k in ("ignored_blank", "ignored_tail", "ignored_non_substantive")
        )
        counts = {
            "read": stats["scanned"],
            "ignored": ignored,
            "imported": imported[spec.name],
            "normalized": normalized[spec.name],
            "rejected": len(rejected),
            "duplicated": len(duplicates),
        }
        counts["accepted"] = counts["normalized"]
        counts["balanced"] = (
            counts["read"] == ignored + counts["imported"]
            and counts["imported"] == counts["normalized"] + counts["rejected"]
        )
        sheets.append(
            {
                "sheet": spec.name,
                "source": spec.logical_source,
                "present": rows is not None,
                **counts,
            }
        )
    return {
        "source": source,
        "status": result["status"],
        "blocking": result["blocking"] or bool(result["blocked_sheets"]),
        "balanced": all(s["balanced"] for s in sheets),
        "contains_source_values": False,
        "sheets": sheets,
        "codes": dict(sorted(Counter(i["code"] for i in result["issues"]).items())),
    }


def golden(report, result):
    # Only synthetic goldens include a semantic digest. Real-file reports never do.
    projection = [
        {
            key: row[key]
            for key in (
                "source",
                "sheet",
                "row",
                "values",
                "original_values",
                "transformations",
            )
        }
        for row in result["normalized_records"]
    ]
    return {**report, "normalized_sha256": digest(canonical(projection))}


def edited(content, *, only=None, action=None):
    book = load_workbook(BytesIO(content))
    try:
        if only:
            for name in list(book.sheetnames):
                if name != only:
                    del book[name]
        if action:
            action(book)
        return _stable_xlsx(book)
    finally:
        book.close()


def corpus_cases():
    clean, _ = workorder(rows=4, tail=2)
    edge, _ = workorder(rows=6, tail=2, edge=True)
    yield "wo-composite", "WO Status", edge
    for name, sheet in (
        ("nfp", "NFP"),
        ("owm", "OWM SN"),
        ("container", "CTNR STATUS"),
        ("blocked", "LOT PENDING_OQC"),
    ):
        yield f"wo-{name}", "WO Status", edited(clean, only=sheet)
    yield (
        "wo-drift",
        "WO Status",
        edited(clean, action=lambda b: setattr(b["NFP"]["Y3"], "value", "SYN-DRIFT")),
    )
    yield "td-clean", "Total Decision", decision()
    yield "td-edge", "Total Decision", decision(edge=True)

    def mixed(book):
        sheet = book.active
        for cell in sheet[4]:
            sheet.cell(5, cell.column, cell.value)
        for cell in sheet[6]:
            cell.value = None
        sheet["A7"] = 46000
        sheet["B8"] = "05/Oct/2026"
        sheet["AX9"] = "#REF!"

    yield "td-mixed", "Total Decision", edited(decision(rows=6), action=mixed)
    yield (
        "td-drift",
        "Total Decision",
        edited(
            decision(), action=lambda b: setattr(b.active["AK1"], "value", "SYN-DRIFT")
        ),
    )

    def unknown(book):
        book.active["BW4"] = "SYNTHETIC-DENIED-CONTENT"
        book.create_sheet("SYNTHETIC-PRIVATE-SHEET")["A1"] = "SYNTHETIC-DENIED-CONTENT"

    yield "td-additional", "Total Decision", edited(decision(), action=unknown)


def verify_manifest(directory=CORPUS):
    manifest = json.loads((directory / "manifest.json").read_text())
    if (
        manifest.get("generator_version") != VERSION
        or manifest.get("schema_version") != VERSION
        or manifest.get("contains_real_data") is not False
        or manifest.get("seed") != 141
    ):
        raise ValueError("Manifesto de ingestão incompatível")
    listed = {"manifest.json", "README.md"}
    ids = set()
    if not isinstance(manifest.get("cases"), list) or not manifest["cases"]:
        raise ValueError("Corpus vazio")
    for case in manifest["cases"]:
        if case["id"] in ids:
            raise ValueError("Caso duplicado")
        ids.add(case["id"])
        specs(case["source"])
        for key in ("input", "golden"):
            name = case[key]
            if Path(name).name != name or name in listed:
                raise ValueError("Caminho de fixture inválido")
            listed.add(name)
            content = (directory / name).read_bytes()
            if digest(content) != case[key + "_sha256"]:
                raise ValueError("Hash de fixture divergente")
    if {p.name for p in directory.iterdir() if p.is_file()} != listed:
        raise ValueError("Inventário do corpus divergente")
    return manifest
