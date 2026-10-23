"""Generate a synthetic workbook with the structure of the WO Status family.

Nothing here comes from the real workbook: only the structure documented in
docs/real-workbook-contracts.md is reproduced (17 sheets, the authorized
sheets' preambles and header rows, mirrored header names backed by formulas,
formula-only columns and trailing rows, hidden sheets) with invented values.
Output is deterministic so tests and operators can compare counts.

    python scripts/generate_wo_status_fixture.py /tmp/wo-status-sintetico.xlsx
    python scripts/generate_wo_status_fixture.py out.xlsx --rows 200 --edge
"""

from __future__ import annotations

import argparse
import sys
from datetime import datetime
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.utils import get_column_letter
from openpyxl.utils.cell import column_index_from_string

ROOT = Path(__file__).resolve().parents[1]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

from scripts.generate_homologation_fixture import _stable_xlsx  # noqa: E402

UNAUTHORIZED_SHEETS = (
    ("Mercosul", "hidden"),
    ("Analysis 2", "visible"),
    ("Analysis NPI", "visible"),
    ("Planilha2", "visible"),
    ("OWMS Hold", "visible"),
    ("RWOK LOCATOR", "visible"),
    ("RWOK LOCATOR (2)", "visible"),
    ("VESSEL DEAD LINE", "visible"),
    ("ICMS", "visible"),
    ("Analysis QTY", "hidden"),
    ("OTD1", "visible"),
    ("Sales", "visible"),
    ("Master", "visible"),
)
SHEET_ORDER = (
    "NFP",
    "Mercosul",
    "OWM SN",
    "LOT PENDING_OQC",
    "Analysis 2",
    "Analysis NPI",
    "Planilha2",
    "OWMS Hold",
    "RWOK LOCATOR",
    "RWOK LOCATOR (2)",
    "VESSEL DEAD LINE",
    "ICMS",
    "Analysis QTY",
    "OTD1",
    "Sales",
    "Master",
    "CTNR STATUS",
)
SENTINEL = "DADO-NAO-AUTORIZADO"


def _put(sheet: Any, letter: str, row: int, value: Any) -> None:
    sheet.cell(row=row, column=column_index_from_string(letter), value=value)


def _headers(sheet: Any, row: int, width: int, named: dict[str, str]) -> None:
    for index in range(1, width + 1):
        letter = get_column_letter(index)
        sheet.cell(row=row, column=index, value=named.get(letter, f"Col {letter}"))


def _nfp(book: Workbook, rows: int, edge: bool) -> dict[str, int]:
    sheet = book["NFP"]
    sheet["A1"] = "Planejamento sintético"
    sheet["A2"] = "Grupo"
    _headers(
        sheet,
        3,
        283,
        {
            "G": "Plant",
            "I": "Demand ID",
            "L": "Lot Qty",
            "M": "OQC Pass",
            "U": "Plant",
            "W": "Line",
            "Y": "Demand ID",
            "AA": "Model",
            "AB": "Suffix",
            "AM": "Lot Qty",
            "HY": "W/O Status",
            "IE": "W/O",
            "JV": "OQC Pass",
        },
    )
    statuses = ("Open", "Closed", "Hold")
    for index in range(rows):
        row = 4 + index
        _put(sheet, "U", row, "SYN")
        _put(sheet, "W", row, f"L{index % 4:02d}")
        _put(sheet, "Y", row, f"D{index:06d}")
        _put(sheet, "AA", row, f"MODEL-{index % 5}")
        _put(sheet, "AB", row, "SFX")
        _put(sheet, "AM", row, 10 + index)
        _put(sheet, "HY", row, statuses[index % 3])
        if index % 10 < 7:
            _put(sheet, "IE", row, f"{index:07d}")
        _put(sheet, "JV", row, "Y" if index % 2 else None)
        _put(sheet, "G", row, f"=U{row}")
        _put(sheet, "I", row, f"=Y{row}")
        _put(sheet, "L", row, f"=AM{row}")
        _put(sheet, "M", row, f'=IF(JV{row}="","",JV{row})')
        _put(sheet, "AO", row, f"=AM{row}*2")
    emitted = rows
    rejected = 0
    ignored_blank = 0
    if edge:
        row = 4 + rows
        _put(sheet, "Y", row, 1234567)  # identifier stored as a number
        _put(sheet, "AM", row, 5)
        row += 1
        _put(sheet, "Y", row, "D-EDGE-NEG")
        _put(sheet, "AM", row, -3)  # negative quantity
        row += 1
        _put(sheet, "Y", row, "D-EDGE-FORMULA")
        _put(sheet, "IE", row, "=A1")  # formula in a mapped column
        row += 1
        _put(sheet, "AA", row, "MODEL-NO-KEY")  # no identity at all
        row += 1
        row += 1  # fully blank row
        ignored_blank += 1
        _put(sheet, "Y", row, "D-EDGE-STATUS")
        _put(sheet, "HY", row, "Totalmente-Desconhecido")  # unknown state
        emitted += 5
        rejected += 3
    return {
        "scanned": rows + (6 if edge else 0),
        "emitted": emitted,
        "rejected_adapter": rejected,
        "ignored_blank": ignored_blank,
    }


def _owm(book: Workbook, rows: int, tail: int, edge: bool) -> dict[str, int]:
    sheet = book["OWM SN"]
    sheet["A1"] = "Recebimento sintético"
    _headers(
        sheet,
        2,
        57,
        {
            "A": "Org",
            "D": "REASON",
            "L": "Model",
            "M": "Pallet Qty",
            "N": "Pallet No",
            "O": "Serial Number",
            "Q": "W/O",
            "R": "OQC Lot",
            "U": "Locator",
            "V": "Status",
            "W": "OQC Pass",
            "Y": "OQC Hold Flag",
        },
    )
    for index in range(rows):
        row = 3 + index
        _put(sheet, "A", row, '="SYN"')
        _put(sheet, "D", row, '=IF(Y3="","","motivo")')
        _put(sheet, "L", row, f"MODEL-{index % 5}")
        _put(sheet, "M", row, 1 + index % 3)
        _put(sheet, "N", row, f"P{index // 4:04d}")
        _put(sheet, "O", row, f"SN{index:08d}")
        _put(sheet, "Q", row, f"{index // 4:07d}")
        _put(sheet, "R", row, f"LOT{index // 8:05d}")
        _put(sheet, "U", row, "A-01")
        _put(sheet, "V", row, "Pending" if index % 3 else "Approved")
        _put(sheet, "W", row, "Y" if index % 2 else "N")
        _put(sheet, "Y", row, "N")
    counts = {
        "scanned": rows + tail,
        "emitted": rows,
        "rejected_adapter": 0,
        "ignored_tail": tail,
        "rejected_gap": 0,
    }
    cursor = 3 + rows
    if edge:
        # a blank gap inside the active region, then two more literal serials
        cursor += 1
        for suffix in ("A", "B"):
            _put(sheet, "O", cursor, f"SN-EDGE-{suffix}")
            _put(sheet, "Q", cursor, "EDGE-WO")
            cursor += 1
        _put(sheet, "O", cursor, "SN-EDGE-A")  # duplicate serial
        _put(sheet, "Q", cursor, "EDGE-WO")
        cursor += 1
        counts["scanned"] += 4
        counts["emitted"] += 3
        counts["rejected_gap"] = 1
    for offset in range(tail):
        row = cursor + offset
        _put(sheet, "A", row, '="SYN"')
        _put(sheet, "D", row, '=IF(Y3="","","motivo")')
    return counts


def _lot(book: Workbook, substantive: int, prefilled: int) -> dict[str, int]:
    sheet = book["LOT PENDING_OQC"]
    sheet["A1"] = "Hold sintético"
    _headers(
        sheet,
        3,
        18,
        {
            "D": "Prod Date",
            "E": "WO",
            "F": "Reason",
            "G": "Hold Group Code",
            "H": "Status",
            "I": "Status Change Date",
            "J": "Request User",
            "K": "Cause Dept.",
            "N": "Change User",
            "O": "Hold Qty",
            "P": "Release Qty",
            "R": "Target Date",
        },
    )
    for index in range(substantive + prefilled):
        row = 4 + index
        _put(sheet, "D", row, '=TODAY()')
        _put(sheet, "E", row, f"=\"WO{index}\"")
        _put(sheet, "F", row, '="motivo"')
        if index < substantive:
            _put(sheet, "G", row, "HG1")
            _put(sheet, "H", row, "Hold")
            _put(sheet, "I", row, datetime(2026, 1, 1 + index))
            _put(sheet, "J", row, "usuario.sintetico")
            _put(sheet, "O", row, 5)
            _put(sheet, "P", row, 0)
    return {
        "scanned": substantive + prefilled,
        "rejected_blocked": substantive,
        "ignored_non_substantive": prefilled,
    }


def _ctnr(book: Workbook, rows: int, edge: bool) -> dict[str, int]:
    sheet = book["CTNR STATUS"]
    _headers(
        sheet,
        1,
        67,
        {
            "C": "CTNR No",
            "D": "CTNR Status",
            "I": "Model",
            "J": "Load Plan Qty",
            "K": "Load Qty",
            "P": "Rmn Qty",
            "AH": "Load Plan Date",
            "AK": "Loading Date",
            "AL": "Load Confirm Date",
            "AW": "Inspection No",
            "AX": "Inspection Status",
            "BE": "W/O",
        },
    )
    for index in range(rows):
        row = 2 + index
        _put(sheet, "C", row, f"CTNR{index:07d}")
        _put(sheet, "D", row, "Loading")
        _put(sheet, "I", row, f"MODEL-{index % 5}")
        _put(sheet, "J", row, 100)
        _put(sheet, "K", row, 40 + index)
        _put(sheet, "P", row, 60 - index)
        _put(sheet, "AH", row, datetime(2026, 2, 1 + index % 20))
        if index % 3:
            _put(sheet, "AK", row, datetime(2026, 2, 2 + index % 20))
        _put(sheet, "BE", row, None)
    counts = {"scanned": rows, "emitted": rows, "rejected_adapter": 0}
    if edge:
        row = 2 + rows
        _put(sheet, "C", row, "CTNR-EDGE-SERIAL")
        _put(sheet, "AH", row, 46000)  # Excel serial date
        _put(sheet, "AK", row, "25/12/2026")  # textual date
        _put(sheet, "AL", row, "amanhã")  # unparseable date, kept with a warning
        row += 1
        _put(sheet, "C", row, "CTNR-EDGE-QTY")
        _put(sheet, "J", row, "abc")  # invalid quantity
        row += 1
        _put(sheet, "C", row, "CTNR-EDGE-TEXT-DATE")
        _put(sheet, "AH", row, "05/Oct/2026 10:30:00")  # month-name text dates
        _put(sheet, "AK", row, "06/Oct/2026")
        counts["scanned"] += 3
        counts["emitted"] += 3
        counts["rejected_adapter"] += 1
    return counts


def build_workbook(
    rows: int = 24, tail: int = 8, edge: bool = False
) -> tuple[bytes, dict[str, dict[str, int]]]:
    """Return the deterministic workbook bytes and the counts it was built with."""
    book = Workbook()
    book.remove(book.active)
    for name in SHEET_ORDER:
        book.create_sheet(name)
    for name, state in UNAUTHORIZED_SHEETS:
        sheet = book[name]
        sheet.append(["Coluna A", "Coluna B"])
        sheet.append([f"{SENTINEL}-{name}", 1])
        sheet.sheet_state = state
    expected = {
        "NFP": _nfp(book, rows, edge),
        "OWM SN": _owm(book, rows, tail, edge),
        "LOT PENDING_OQC": _lot(book, substantive=6, prefilled=10),
        "CTNR STATUS": _ctnr(book, max(rows // 2, 1), edge),
    }
    return _stable_xlsx(book), expected


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    parser.add_argument("output", type=Path)
    parser.add_argument("--rows", type=int, default=24)
    parser.add_argument("--tail", type=int, default=8)
    parser.add_argument("--edge", action="store_true")
    arguments = parser.parse_args()
    content, expected = build_workbook(arguments.rows, arguments.tail, arguments.edge)
    arguments.output.write_bytes(content)
    print(f"{arguments.output}: {len(content)} bytes; contagens esperadas: {expected}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
